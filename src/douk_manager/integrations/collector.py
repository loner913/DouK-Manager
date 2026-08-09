from __future__ import annotations

import os
import hashlib
import shutil
import subprocess
import sys
import time
import urllib.request
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from douk_manager.config import AppConfig, ManagedPaths, project_root, resource_path
from douk_manager.core.json_store import read_json


class CollectorServiceError(RuntimeError):
    pass


@dataclass(frozen=True)
class MigrationResult:
    copied_files: tuple[str, ...]
    skipped_files: tuple[str, ...]
    copied_screenshots: int
    skipped_screenshots: int
    excel_original_max: int
    excel_final_max: int
    excel_rows_added: int
    excel_existing_url_differences: int


class CollectorService:
    STARTUP_TIMEOUT_SECONDS = 15.0

    def __init__(self, config: AppConfig, paths: ManagedPaths) -> None:
        self.config = config
        self.paths = paths
        self.process: subprocess.Popen | None = None
        self._log_handle = None
        self.last_log_path: Path | None = None

    @property
    def running(self) -> bool:
        return self.process is not None and self.process.poll() is None

    def health(self, timeout: float = 0.8) -> bool:
        try:
            with urllib.request.urlopen(
                f"http://127.0.0.1:{self.config.collector_port}/health",
                timeout=timeout,
            ) as response:
                return response.status == 200
        except Exception:
            return False

    def start(self) -> None:
        if self.running:
            raise CollectorServiceError("账号采集服务已经在运行。")
        if self.process is not None:
            self.process = None
            self._close_log_handle()
        if self.health():
            raise CollectorServiceError(
                f"端口 {self.config.collector_port} 已有采集服务在运行，无需重复启动。"
            )
        if not self.paths.master_settings.is_file():
            raise CollectorServiceError(f"唯一正式主档不存在：{self.paths.master_settings}")
        if not self.paths.collector_excel.is_file():
            raise CollectorServiceError(
                f"采集器 Excel 不存在：{self.paths.collector_excel}。请先迁移旧采集器数据。"
            )
        self.paths.collector_data.mkdir(parents=True, exist_ok=True)
        self.paths.screenshot_inbox.mkdir(parents=True, exist_ok=True)
        env = os.environ.copy()
        env.update(
            {
                "DOUK_COLLECTOR_DATA_DIR": str(self.paths.collector_data),
                "DOUK_MASTER_PATH": str(self.paths.master_settings),
                "DOUK_COLLECTOR_EXCEL": str(self.paths.collector_excel),
                "DOUK_SCREENSHOT_DIR": str(self.paths.screenshot_inbox),
                "DOUK_COLLECTOR_PORT": str(self.config.collector_port),
                "DOUK_COLLECTOR_TOKEN": self.config.collector_token,
                "DOUK_GLOBAL_LOCK_PATH": str(self.paths.lock_file),
                "PYTHONUNBUFFERED": "1",
                "PYTHONIOENCODING": "utf-8",
                "PYTHONUTF8": "1",
            }
        )
        if getattr(sys, "frozen", False):
            command = [sys.executable, "--collector-worker"]
        else:
            command = [sys.executable, str(project_root() / "main.py"), "--collector-worker"]
            existing = env.get("PYTHONPATH", "")
            env["PYTHONPATH"] = os.pathsep.join(
                part for part in (str(project_root() / "src"), existing) if part
            )
        log_path = self.paths.logs / f"Collector_{datetime.now():%Y-%m-%d_%H-%M-%S}.log"
        self.last_log_path = log_path
        self._log_handle = log_path.open("a", encoding="utf-8")
        creationflags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
        try:
            self.process = subprocess.Popen(
                command,
                cwd=str(self.paths.root),
                env=env,
                stdout=self._log_handle,
                stderr=subprocess.STDOUT,
                creationflags=creationflags,
            )
        except OSError as exc:
            self._close_log_handle()
            raise CollectorServiceError(f"无法启动账号采集服务：{exc}") from exc
        self._wait_until_ready(log_path)

    def _wait_until_ready(self, log_path: Path) -> None:
        """Only report success after the HTTP service is genuinely reachable."""
        deadline = time.monotonic() + self.STARTUP_TIMEOUT_SECONDS
        while True:
            process = self.process
            if process is None:
                raise CollectorServiceError("账号采集服务进程状态丢失，启动已取消。")
            exit_code = process.poll()
            if exit_code is not None:
                self.process = None
                self._close_log_handle()
                details = _read_log_tail(log_path)
                raise CollectorServiceError(
                    "账号采集服务未能启动"
                    f"（进程退出码 {exit_code}）。\n"
                    f"日志：{log_path}\n"
                    f"{details or '日志中没有输出。'}"
                )
            if self.health(timeout=0.35):
                return
            if time.monotonic() >= deadline:
                self.stop()
                details = _read_log_tail(log_path)
                raise CollectorServiceError(
                    "账号采集服务进程已创建，但在 15 秒内没有成功监听 "
                    f"127.0.0.1:{self.config.collector_port}，已自动停止。\n"
                    f"日志：{log_path}\n"
                    f"{details or '日志中没有输出。'}"
                )
            time.sleep(0.15)

    def stop(self) -> None:
        if self.process is not None and self.process.poll() is None:
            self.process.terminate()
            try:
                self.process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                self.process.kill()
                self.process.wait(timeout=3)
        self.process = None
        self._close_log_handle()

    def _close_log_handle(self) -> None:
        if self._log_handle is not None:
            try:
                self._log_handle.flush()
            finally:
                self._log_handle.close()
                self._log_handle = None

    def migrate_old_data(self, old_screenshot_dir: Path) -> MigrationResult:
        old_root = old_screenshot_dir.parent
        self.paths.collector_data.mkdir(parents=True, exist_ok=True)
        self.paths.screenshot_inbox.mkdir(parents=True, exist_ok=True)
        copied: list[str] = []
        skipped: list[str] = []
        for filename in (
            "录制名单.xlsx",
            ".douk_backup_state.json",
            "顶级.txt",
            "次顶级.txt",
            "普通.txt",
        ):
            source = old_root / filename
            target = self.paths.collector_data / filename
            if not source.is_file():
                skipped.append(filename + "（旧文件不存在）")
            elif target.exists():
                skipped.append(filename + "（新位置已有文件）")
            else:
                _copy_verified_exclusive(source, target)
                copied.append(filename)
        copied_screenshots = 0
        skipped_screenshots = 0
        if old_screenshot_dir.is_dir():
            for source in old_screenshot_dir.iterdir():
                if not source.is_file():
                    continue
                target = self.paths.screenshot_inbox / source.name
                if target.exists():
                    skipped_screenshots += 1
                    continue
                _copy_verified_exclusive(source, target)
                copied_screenshots += 1
        reconcile = reconcile_excel_to_master(
            self.paths.collector_excel,
            self.paths.master_settings,
            self.paths.backups / "CollectorMigration",
        )
        return MigrationResult(
            tuple(copied),
            tuple(skipped),
            copied_screenshots,
            skipped_screenshots,
            reconcile.original_max,
            reconcile.final_max,
            reconcile.rows_added,
            reconcile.existing_url_differences,
        )

    def export_userscript(self) -> Path:
        source = resource_path("resources/userscript/douyin_account_collector.user.js")
        target = self.paths.data / "douyin_account_collector.user.js"
        shutil.copy2(source, target)
        return target


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _read_log_tail(path: Path, max_lines: int = 30, max_chars: int = 6000) -> str:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError:
        return ""
    tail = "\n".join(text.splitlines()[-max_lines:])
    return tail[-max_chars:]


def _copy_verified_exclusive(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_name(f".{target.name}.migrating")
    if target.exists() or temp.exists():
        raise CollectorServiceError(f"迁移目标已存在，拒绝覆盖：{target}")
    try:
        shutil.copy2(source, temp)
        if source.stat().st_size != temp.stat().st_size or _sha256(source) != _sha256(temp):
            raise CollectorServiceError(f"迁移副本校验失败：{source}")
        os.replace(temp, target)
    except Exception:
        temp.unlink(missing_ok=True)
        raise


@dataclass(frozen=True)
class ExcelReconcileResult:
    original_max: int
    final_max: int
    rows_added: int
    backup_path: Path | None
    existing_url_differences: int


def reconcile_excel_to_master(
    excel_path: Path, master_path: Path, backup_root: Path
) -> ExcelReconcileResult:
    """将 Excel 的尾部缺失记录补齐到唯一主档；不改已有行、不改主档。"""
    if not excel_path.is_file():
        raise CollectorServiceError(f"无法对齐，Excel 不存在：{excel_path}")
    master = read_json(master_path)
    accounts = master.get("accounts_urls")
    if not isinstance(accounts, list):
        raise CollectorServiceError("唯一主档缺少 accounts_urls 数组。")

    occupied: list[tuple[int, dict]] = []
    encountered_blank = False
    for number, account in enumerate(accounts, start=1):
        if not isinstance(account, dict):
            raise CollectorServiceError(f"唯一主档 A{number} 不是对象。")
        mark = str(account.get("mark", "")).strip()
        url = str(account.get("url", "")).strip()
        if bool(mark) != bool(url):
            raise CollectorServiceError(f"唯一主档 A{number} 是 mark/url 半条记录。")
        if not mark:
            encountered_blank = True
            continue
        if encountered_blank:
            raise CollectorServiceError(
                f"唯一主档 A{number} 位于空白位置之后，编号不连续，拒绝自动对齐Excel。"
            )
        prefix = f"A{number}"
        if not mark.casefold().startswith(prefix.casefold()) or not mark[len(prefix) :].strip():
            raise CollectorServiceError(f"唯一主档 A{number} 的 mark 与数组位置不一致。")
        occupied.append((number, account))
    master_max = occupied[-1][0] if occupied else 0

    workbook = load_workbook(excel_path, data_only=False, read_only=False, keep_links=True)
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise CollectorServiceError("录制名单.xlsx 缺少 Sheet1。")
        sheet = workbook["Sheet1"]
        excel_numbers: list[int] = []
        for row in range(2, 5002):
            b_value = str(sheet[f"B{row}"].value or "").strip()
            d_value = str(sheet[f"D{row}"].value or "").strip()
            link = sheet[f"D{row}"].hyperlink
            if bool(b_value) != bool(d_value):
                raise CollectorServiceError(f"Excel A{row - 1} 存在 B/D 半行记录。")
            if not d_value and link:
                raise CollectorServiceError(f"Excel A{row - 1} 存在残留超链接。")
            if b_value:
                excel_numbers.append(row - 1)
        excel_max = excel_numbers[-1] if excel_numbers else 0
        if excel_numbers != list(range(1, excel_max + 1)):
            raise CollectorServiceError("Excel现有账号存在中间空行，拒绝自动修复。")
        if excel_max > master_max:
            raise CollectorServiceError(
                f"Excel 已到 A{excel_max}，超过主档 A{master_max}，拒绝自动覆盖。"
            )

        # 旧 Excel 可能保存短链、历史链接或不同写法。它们都是既有采集数据，
        # 只统计差异并原样保留；自动对齐仅允许补齐 Excel 尾部缺失账号。
        existing_url_differences = 0
        for number in range(1, excel_max + 1):
            excel_url = str(sheet[f"D{number + 1}"].value or "").strip()
            master_url = str(accounts[number - 1].get("url", "")).strip()
            if excel_url != master_url:
                existing_url_differences += 1

        if excel_max == master_max:
            return ExcelReconcileResult(
                excel_max, master_max, 0, None, existing_url_differences
            )

        stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S-%f")
        backup_dir = backup_root / stamp
        backup_dir.mkdir(parents=True, exist_ok=False)
        backup_path = backup_dir / excel_path.name
        _copy_verified_exclusive(excel_path, backup_path)

        merged = {str(item) for item in sheet.merged_cells.ranges}
        for number in range(excel_max + 1, master_max + 1):
            row = number + 1
            if f"B{row}:C{row}" not in merged or f"D{row}:J{row}" not in merged:
                raise CollectorServiceError(f"Excel A{number} 模板合并单元格不完整。")
            account = accounts[number - 1]
            mark = str(account.get("mark", "")).strip()
            url = str(account.get("url", "")).strip()
            prefix = f"A{number}"
            suffix = mark[len(prefix) :].strip()
            sheet[f"B{row}"].value = suffix
            sheet[f"D{row}"].hyperlink = None
            sheet[f"D{row}"].value = url

        temp = excel_path.with_name(f".{excel_path.name}.reconcile.tmp.xlsx")
        try:
            workbook.save(temp)
            verify = load_workbook(temp, data_only=False, read_only=False, keep_links=True)
            try:
                verify_sheet = verify["Sheet1"]
                # 已有行必须逐字保留；新增尾部行才应与主档一致。
                for number in range(1, excel_max + 1):
                    expected_url = str(sheet[f"D{number + 1}"].value or "").strip()
                    actual_url = str(verify_sheet[f"D{number + 1}"].value or "").strip()
                    if actual_url != expected_url:
                        raise CollectorServiceError(
                            f"Excel临时副本 A{number} 历史行发生变化，拒绝替换。"
                        )
                for number in range(excel_max + 1, master_max + 1):
                    account = accounts[number - 1]
                    expected_url = str(account.get("url", "")).strip()
                    actual_url = str(verify_sheet[f"D{number + 1}"].value or "").strip()
                    if actual_url != expected_url:
                        raise CollectorServiceError(f"Excel临时副本 A{number} 复读校验失败。")
            finally:
                verify.close()
            os.replace(temp, excel_path)
        except Exception:
            temp.unlink(missing_ok=True)
            raise
        return ExcelReconcileResult(
            excel_max,
            master_max,
            master_max - excel_max,
            backup_path,
            existing_url_differences,
        )
    finally:
        workbook.close()
