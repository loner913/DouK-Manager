from __future__ import annotations

import http.client
import json
import os
import socket
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import uuid
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch
from urllib.request import urlopen

from openpyxl import Workbook

from douk_manager.core.watchlist import WatchlistService
from douk_manager.core import watchlist as watchlist_module
from douk_manager.core.watchlist_session import ManagerInstanceLease
from douk_manager.vendor import collector_server
from tests.helpers import make_test_paths


class CollectorWatchlistHttpTests(unittest.TestCase):
    def test_active_request_is_busy_and_abandoned_pending_requires_recovery(self) -> None:
        with tempfile.TemporaryDirectory() as directory, ExitStack() as stack:
            paths = make_test_paths(Path(directory), account_count=2)
            service = WatchlistService(paths)
            service.initialize(initialization_evidence=True)
            service.mark_ready()
            for name, value in {
                "SETTINGS_PATH": paths.master_settings,
                "WATCHLIST_PATH": paths.watchlist,
                "WATCHLIST_WATERMARK_PATH": paths.watchlist_w_watermark,
                "WATCHLIST_CONTROL_PATH": paths.watchlist_control,
                "GLOBAL_LOCK_PATH": paths.lock_file,
                "MANAGER_INSTANCE_LOCK_PATH": paths.instance_lock_file,
            }.items():
                stack.enter_context(patch.object(collector_server, name, value))
            for name in ("manager_session_active", "manager_instance_owner_present"):
                stack.enter_context(patch.object(collector_server, name, return_value=True))
            server = collector_server.ThreadingHTTPServer(("127.0.0.1", 0), collector_server.Handler)
            server_thread = threading.Thread(target=server.serve_forever)
            server_thread.start()
            entered, release = threading.Event(), threading.Event()
            original = self._payload()
            other = {**original, "request_id": str(uuid.uuid4()), "note": "other request"}
            real_write = watchlist_module.write_json_atomic
            failures = []

            def interrupt_body(path, value):
                if path == paths.watchlist:
                    entered.set()
                    if not release.wait(5):
                        raise RuntimeError("synthetic request was not released")
                    raise PermissionError("synthetic interrupted body")
                real_write(path, value)

            def original_request():
                try:
                    service.observe(original)
                except Exception as exc:
                    failures.append(exc)

            def post(payload):
                connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=2)
                try:
                    connection.request("POST", "/watchlist/observe", json.dumps(payload), {
                        "Content-Type": "application/json", "X-DouK-Token": collector_server.ACCESS_TOKEN,
                    })
                    response = connection.getresponse()
                    return response.status, json.loads(response.read())
                finally:
                    connection.close()

            request_thread = threading.Thread(target=original_request)
            try:
                with patch.object(watchlist_module, "write_json_atomic", side_effect=interrupt_body):
                    request_thread.start()
                    try:
                        self.assertTrue(entered.wait(2))
                        targets = (paths.watchlist, paths.watchlist_control, paths.watchlist_w_watermark)
                        before = [path.read_bytes() for path in targets]
                        status, body = post(other)
                        self.assertEqual((status, body["code"], body["details"]["message_code"]), (409, "BUSY", "BUSY"))
                        self.assertEqual(before, [path.read_bytes() for path in targets])
                    finally:
                        release.set()
                        request_thread.join(5)
                self.assertFalse(request_thread.is_alive())
                self.assertEqual(len(failures), 1)
                self.assertIsInstance(failures[0], PermissionError)
                status, body = post(other)
                self.assertEqual((status, body["code"]), (503, "RECOVERY_REQUIRED"))
                self.assertEqual(before, [path.read_bytes() for path in targets])
                status, body = post(original)
                self.assertEqual((status, body["status"], body["w_id"]), (200, "EXISTS", 1))
                self.assertEqual(service.snapshot().next_w_id, 2)
            finally:
                release.set()
                if request_thread.ident is not None:
                    request_thread.join(5)
                server.shutdown()
                server.server_close()
                server_thread.join(5)

    @staticmethod
    def _payload() -> dict[str, object]:
        return {
            "request_id": str(uuid.uuid4()),
            "url": "https://www.douyin.com/user/42",
            "captured_nickname": "合成昵称",
            "display_name": "合成名称",
            "douyin_id": "synthetic_42",
            "nickname_blank": False,
            "reasons": ["few_works"],
            "note": "synthetic",
            "review_after_days": 7,
            "client_version": collector_server.VERSION,
            "page_visible": True,
        }

    def test_observe_requires_session_but_formal_server_shell_stays_available(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            paths = make_test_paths(Path(directory), account_count=2)
            WatchlistService(paths).initialize(initialization_evidence=True)
            WatchlistService(paths).mark_ready()
            old_values = {
                "SETTINGS_PATH": collector_server.SETTINGS_PATH,
                "WATCHLIST_PATH": collector_server.WATCHLIST_PATH,
                "WATCHLIST_WATERMARK_PATH": collector_server.WATCHLIST_WATERMARK_PATH,
                "WATCHLIST_CONTROL_PATH": collector_server.WATCHLIST_CONTROL_PATH,
                "GLOBAL_LOCK_PATH": collector_server.GLOBAL_LOCK_PATH,
                "MANAGER_INSTANCE_LOCK_PATH": collector_server.MANAGER_INSTANCE_LOCK_PATH,
            }
            collector_server.SETTINGS_PATH = paths.master_settings
            collector_server.WATCHLIST_PATH = paths.watchlist
            collector_server.WATCHLIST_WATERMARK_PATH = paths.watchlist_w_watermark
            collector_server.WATCHLIST_CONTROL_PATH = paths.watchlist_control
            collector_server.GLOBAL_LOCK_PATH = paths.lock_file
            collector_server.MANAGER_INSTANCE_LOCK_PATH = paths.instance_lock_file
            server = collector_server.ThreadingHTTPServer(("127.0.0.1", 0), collector_server.Handler)
            worker = threading.Thread(target=server.serve_forever, daemon=True)
            worker.start()
            try:
                body = json.dumps(self._payload(), ensure_ascii=False).encode("utf-8")
                connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=3)
                connection.request(
                    "POST",
                    "/watchlist/observe",
                    body=body,
                    headers={
                        "Content-Type": "application/json",
                        "Content-Length": str(len(body)),
                        "X-DouK-Token": collector_server.ACCESS_TOKEN,
                    },
                )
                response = connection.getresponse()
                blocked = json.loads(response.read().decode("utf-8"))
                connection.close()
                self.assertEqual(response.status, 503)
                self.assertEqual(blocked["code"], "MANAGER_SESSION_REQUIRED")
                self.assertEqual(read_records(paths.watchlist), [])

                lease = ManagerInstanceLease(paths.instance_lock_file)
                lease.acquire()
                try:
                    with patch.object(collector_server, "manager_session_active", return_value=True), patch.object(
                        collector_server, "manager_instance_owner_present", return_value=True
                    ):
                        body = json.dumps(self._payload(), ensure_ascii=False).encode("utf-8")
                        connection = http.client.HTTPConnection(
                            "127.0.0.1", server.server_port, timeout=3
                        )
                        connection.request(
                            "POST",
                            "/watchlist/observe",
                            body=body,
                            headers={
                                "Content-Type": "application/json",
                                "Content-Length": str(len(body)),
                                "X-DouK-Token": collector_server.ACCESS_TOKEN,
                            },
                        )
                        response = connection.getresponse()
                        created = json.loads(response.read().decode("utf-8"))
                        connection.close()
                    self.assertEqual(response.status, 201)
                    self.assertEqual(created["status"], "CREATED")
                    self.assertEqual(created["w_id"], 1)
                    self.assertNotIn("url", created)
                    original_payload = json.loads(body)
                    watchlist = WatchlistService(paths)
                    archived = watchlist.archive(1, expected_revision=watchlist.snapshot().revision)
                    watchlist.delete_archived(1, request_id=str(uuid.uuid4()), expected_revision=archived.revision)
                    before = [path.read_bytes() for path in (paths.watchlist, paths.watchlist_control, paths.watchlist_w_watermark)]
                    with patch.object(collector_server, "manager_session_active", return_value=True):
                        connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=3)
                        try:
                            connection.request("POST", "/watchlist/observe", json.dumps(original_payload),
                                               {"Content-Type": "application/json", "X-DouK-Token": collector_server.ACCESS_TOKEN})
                            response = connection.getresponse()
                            replayed = json.loads(response.read())
                            self.assertEqual(response.status, 410)
                            self.assertEqual((replayed["status"], replayed["w_id"]), ("DELETED", 1))
                        finally:
                            connection.close()
                    self.assertEqual(before, [path.read_bytes() for path in (paths.watchlist, paths.watchlist_control, paths.watchlist_w_watermark)])
                finally:
                    lease.release()
            finally:
                server.shutdown()
                server.server_close()
                worker.join(timeout=2)
                for name, value in old_values.items():
                    setattr(collector_server, name, value)

    def test_real_worker_pipe_eof_exits_without_further_observation_writes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            paths = make_test_paths(root, account_count=2)
            WatchlistService(paths).initialize(initialization_evidence=True)
            WatchlistService(paths).mark_ready()
            create_synthetic_collector_workbook(paths.collector_excel)
            port = free_port()
            environment = dict(os.environ)
            environment.update(
                {
                    "DOUK_COLLECTOR_DATA_DIR": str(paths.collector_data),
                    "DOUK_MASTER_PATH": str(paths.master_settings),
                    "DOUK_COLLECTOR_EXCEL": str(paths.collector_excel),
                    "DOUK_SCREENSHOT_DIR": str(paths.screenshot_inbox),
                    "DOUK_COLLECTOR_PORT": str(port),
                    "DOUK_COLLECTOR_TOKEN": "synthetic-token",
                    "DOUK_GLOBAL_LOCK_PATH": str(paths.lock_file),
                    "DOUK_MANAGER_SESSION_CHANNEL": "1",
                    "DOUK_WATCHLIST_PATH": str(paths.watchlist),
                    "DOUK_WATCHLIST_WATERMARK_PATH": str(paths.watchlist_w_watermark),
                    "DOUK_WATCHLIST_CONTROL_PATH": str(paths.watchlist_control),
                    "DOUK_MANAGER_INSTANCE_LOCK_PATH": str(paths.instance_lock_file),
                    "PYTHONUNBUFFERED": "1",
                    "PYTHONIOENCODING": "utf-8",
                    "PYTHONUTF8": "1",
                }
            )
            lease = ManagerInstanceLease(paths.instance_lock_file)
            lease.acquire()
            process = None
            try:
                process = subprocess.Popen(
                    [sys.executable, str(Path(__file__).resolve().parents[1] / "main.py"), "--collector-worker"],
                    cwd=str(Path(__file__).resolve().parents[1]),
                    env=environment,
                    stdin=subprocess.PIPE,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                )
                assert process.stdin is not None
                process.stdin.write(b"synthetic-manager-session\n")
                process.stdin.flush()
                wait_for_health(port, "synthetic-token")
                created = post_observe(port, "synthetic-token", self._payload())
                self.assertEqual(created["status"], "CREATED")
                self.assertEqual(created["w_id"], 1)

                # This explicit close simulates the manager process exiting;
                # production CollectorService never closes its write end
                # during normal operation.
                before = {
                    path: path.read_bytes()
                    for path in (paths.watchlist, paths.watchlist_control, paths.watchlist_w_watermark)
                }
                process.stdin.close()
                process.stdin = None
                output, _ = process.communicate(timeout=5)
                self.assertEqual(process.returncode, 0, output.decode("utf-8", errors="replace"))
                for path, content in before.items():
                    self.assertEqual(path.read_bytes(), content)
                with self.assertRaises((ConnectionError, OSError)):
                    post_observe(port, "synthetic-token", self._payload())
            finally:
                try:
                    if process is not None and process.stdin is not None:
                        try:
                            process.stdin.close()
                        except (OSError, ValueError):
                            pass
                    if process is not None and process.poll() is None:
                        process.terminate()
                    if process is not None:
                        try:
                            process.communicate(timeout=5)
                        except subprocess.TimeoutExpired:
                            process.kill()
                            process.communicate(timeout=5)
                finally:
                    lease.release()


def create_synthetic_collector_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = collector_server.EXCEL_SHEET_NAME
    sheet["A2"] = "1.A1"
    sheet["B2"] = "account1"
    sheet["D2"] = "https://www.douyin.com/user/1"
    sheet["A3"] = "2.A2"
    sheet["B3"] = "account2"
    sheet["D3"] = "https://www.douyin.com/user/2"
    sheet["A4"] = "3.A3"
    sheet.merge_cells("B4:C4")
    sheet.merge_cells("D4:J4")
    workbook.save(path)
    workbook.close()


def free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return int(sock.getsockname()[1])


def wait_for_health(port: int, token: str) -> None:
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        try:
            with urlopen(
                f"http://127.0.0.1:{port}/health",
                timeout=0.25,
            ) as response:
                if response.status == 200:
                    return
        except OSError:
            time.sleep(0.05)
    raise AssertionError("synthetic collector worker did not become healthy")


def post_observe(port: int, token: str, payload: dict[str, object]) -> dict[str, object]:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    connection = http.client.HTTPConnection("127.0.0.1", port, timeout=2)
    try:
        connection.request(
            "POST",
            "/watchlist/observe",
            body=body,
            headers={
                "Content-Type": "application/json",
                "Content-Length": str(len(body)),
                "X-DouK-Token": token,
            },
        )
        response = connection.getresponse()
        return json.loads(response.read().decode("utf-8"))
    finally:
        connection.close()


def read_records(path: Path) -> list[dict[str, object]]:
    return json.loads(path.read_text(encoding="utf-8"))["records"]


if __name__ == "__main__":
    unittest.main()
