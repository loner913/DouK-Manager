from __future__ import annotations

import ast
import hashlib
import importlib.util
import inspect
import io
import json
import os
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unittest
from contextlib import closing
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from douk_manager.background import (
    CancellationToken,
    TaskFailure,
    TaskRejectedError,
    TaskState,
    TaskWorker,
)
from douk_manager.config import AppConfig
from douk_manager.controller import ControllerError, ManagerController
from douk_manager.core.backup import BackupService
from douk_manager.core.engine import (
    EngineError,
    EngineService,
    ProcessProbe,
    ProcessProbeState,
    _WindowsEngineMutex,
)
from douk_manager.core.json_store import read_json, write_json_atomic
from douk_manager.core.settings_tasks import EarliestRule, SettingsTaskService
from douk_manager.gui import MainWindow
from douk_manager.integrations import indexer as indexer_module
from douk_manager.integrations.indexer import IndexService
from douk_manager.integrations.screenshots import ScreenshotService
from douk_manager.operation import OperationContext, TaskCancelled
from douk_manager.vendor import collector_server, screenshot_organizer
from tests.helpers import make_test_paths

try:
    from douk_manager import startup as startup_module
except ImportError:
    startup_module = None


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def index_service_output(mode: str) -> str:
    summary: dict[str, object] = {
        "Mode": mode,
        "SourceRoot": "synthetic-source",
        "IndexRoot": "synthetic-index",
        "SourceFoldersScanned": 1,
        "IgnoredSourceFolders": 0,
        "EmptySourceFolders": 0,
        "MovedOrDeletedTargetFolders": 0,
        "EmptySourceShortcutsDetected": 0,
        "MissingTargetShortcutsDetected": 0,
        "PlannedShortcutDeletions": 0,
        "DeletedEmptySourceShortcuts": 0,
        "DeletedMissingTargetShortcuts": 0,
        "DeletedShortcutsTotal": 0,
        "ShortcutDeleteFailures": 0,
        "RemainingEmptySourceShortcuts": 0,
        "RemainingMissingTargetShortcuts": 0,
        "ShortcutReadFailures": 0,
    }
    if mode == "Refresh":
        summary.update(Created=5, Updated=6, Unchanged=7, IndexFailures=0)
    return "synthetic index details\nDOUK_INDEX_SUMMARY_JSON=" + json.dumps(summary)


def make_controller(root: Path, *, engine_running: bool) -> ManagerController:
    paths = make_test_paths(root, account_count=3)
    controller = ManagerController.__new__(ManagerController)
    controller.paths = paths
    controller.config = SimpleNamespace(collector_port=18765)
    controller.logger = Mock()
    controller.startup_backup = paths.backups / "existing-startup-backup"
    controller.read_only_reason = ""
    controller._last_collector_running = False
    controller._last_engine_running = engine_running
    controller.engine = Mock()
    controller.engine.external_running.return_value = engine_running
    controller.collector = Mock()
    controller.collector.last_log_path = paths.logs / "Collector_test.log"
    controller.tasks = Mock()
    controller.backup = Mock()
    controller.engine_updates = Mock()
    return controller


def create_collector_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = collector_server.EXCEL_SHEET_NAME
    sheet["A2"] = "1.A1"
    sheet["B2"] = "existing111"
    sheet["D2"] = "https://www.douyin.com/user/existing"
    sheet["A3"] = "2.A2"
    sheet.merge_cells("B3:C3")
    sheet.merge_cells("D3:J3")
    workbook.save(path)
    workbook.close()


class ControllerRuntimeSafetyTests(unittest.TestCase):
    def _startup_state(self):
        self.assertIsNotNone(
            startup_module,
            "Phase 1 must provide the douk_manager.startup state contracts.",
        )
        assert startup_module is not None
        self.assertTrue(
            hasattr(startup_module, "StartupState"),
            "Phase 1 must expose douk_manager.startup.StartupState.",
        )
        return startup_module.StartupState

    def _degraded_controller(self, root: Path) -> ManagerController:
        state = self._startup_state()
        controller = make_controller(root, engine_running=False)
        controller.root = controller.paths.root
        controller.config = AppConfig(
            engine_exe=str(controller.paths.engine_exe),
            video_root=str(controller.paths.video_root),
            index_root=str(controller.paths.index_root),
        )
        controller.task_order = Mock()
        controller.screenshots = Mock()
        controller.indexer = Mock()
        controller.startup_state = state.DEGRADED_READ_ONLY
        controller.startup_generation = 4
        controller.read_only_reason = "synthetic startup safety failure"
        controller.collector.running = False
        controller.collector.health.return_value = False
        controller.engine.current = None
        controller.engine.external_running.return_value = False
        return controller

    def _post_action_controller(
        self,
        *,
        screenshot_mode: str = "batch",
        index_mode: str = "batch",
        cleanup: bool = True,
        lifecycle_active: bool = True,
    ) -> ManagerController:
        controller = ManagerController.__new__(ManagerController)
        controller.startup_state = self._startup_state().READY
        controller.read_only_reason = ""
        controller._download_lifecycle_active = lifecycle_active
        controller.config = SimpleNamespace(
            screenshot_post_mode=screenshot_mode,
            index_post_mode=index_mode,
            cleanup_after_index=cleanup,
        )
        controller.paths = SimpleNamespace(
            screenshot_inbox=Path("synthetic-screenshots"),
            video_root=Path("synthetic-videos"),
            index_root=Path("synthetic-index"),
            index_refresh_logs=Path("synthetic-refresh-logs"),
            index_cleanup_logs=Path("synthetic-cleanup-logs"),
        )
        controller.logger = Mock()
        controller.screenshots = Mock()
        controller.screenshots.execute.return_value = SimpleNamespace(moved=2)
        controller.indexer = Mock()
        refresh_result = Mock(name="refresh_result")
        refresh_result.display_lines.return_value = ("refresh numeric log",)
        refresh_result.display_summary.return_value = "refresh numeric summary"
        cleanup_result = Mock(name="cleanup_result")
        cleanup_result.display_lines.return_value = ("cleanup numeric log",)
        cleanup_result.display_summary.return_value = "cleanup numeric summary"
        controller.indexer.refresh.return_value = refresh_result
        controller.indexer.cleanup.return_value = cleanup_result
        return controller

    @staticmethod
    def _create_screenshot_inputs(
        root: Path,
        *,
        count: int = 1,
    ) -> tuple[Path, Path, tuple[Path, ...], tuple[Path, ...]]:
        inbox = root / "screenshots"
        accounts = root / "accounts"
        inbox.mkdir()
        accounts.mkdir()
        sources: list[Path] = []
        destinations: list[Path] = []
        for number in range(1, count + 1):
            destination = accounts / f"UID100{number}_A{number}account_works"
            destination.mkdir()
            source = inbox / f"A{number}.jpg"
            source.write_bytes(f"screenshot-{number}".encode("ascii"))
            sources.append(source)
            destinations.append(destination / source.name)
        return inbox, accounts, tuple(sources), tuple(destinations)

    def _read_only_controller(self, state) -> ManagerController:
        controller = ManagerController.__new__(ManagerController)
        controller.startup_state = state
        controller.read_only_reason = "synthetic startup gate"
        controller.logger = Mock()
        controller.results = Mock()
        controller.results.page_snapshot.return_value = "result-snapshot"
        controller.results.classify_private_reference.return_value = ()
        controller.task_order = Mock()
        controller.task_order.list_tasks.return_value = (Path("Task_A1.json"),)
        controller.task_order.last_warning = ""
        controller.tasks = Mock()
        controller.tasks.preview.return_value = SimpleNamespace(
            selection=SimpleNamespace(numbers=())
        )
        controller.tasks.preview_with_private_filter.return_value = "private-preview"
        controller.screenshots = Mock()
        controller.screenshots.preview.return_value = "screenshot-preview"
        controller.paths = SimpleNamespace(
            screenshot_inbox=Path("synthetic-screenshots"),
            video_root=Path("synthetic-videos"),
        )
        return controller

    def _runtime_diagnostic_controller(self, state=None) -> ManagerController:
        state = state or self._startup_state().READY
        controller = ManagerController.__new__(ManagerController)
        controller.startup_state = state
        controller.read_only_reason = "synthetic startup gate"
        controller.startup_backup = None
        controller._last_collector_running = False
        controller._last_engine_running = False
        controller.logger = Mock()
        controller.paths = SimpleNamespace(
            health=Mock(return_value={"engine_exe": True}),
            active_settings=Mock(is_file=Mock(return_value=False)),
            master_settings=Mock(is_file=Mock(return_value=False)),
        )
        controller.collector = SimpleNamespace(
            health=Mock(return_value=False),
            process=None,
            running=False,
        )
        controller.engine = SimpleNamespace(
            external_running=Mock(return_value=False),
            current=None,
        )
        return controller

    @staticmethod
    def _invoke_read_only_controller(
        controller: ManagerController,
        entry: str,
        *,
        context: object | None = None,
    ):
        if entry == "result":
            return controller.result_snapshot(limit=25, context=context)
        if entry == "task":
            return controller.list_tasks()
        if entry == "private":
            return controller.preview_private_skip("A1", 30, context=context)
        if entry == "screenshot":
            return controller.screenshot_preview(context=context)
        raise AssertionError(entry)

    def test_read_only_controller_entries_have_independent_ready_second_gate(
        self,
    ) -> None:
        state_type = self._startup_state()
        expected = {
            "result": "result-snapshot",
            "task": (Path("Task_A1.json"),),
            "private": "private-preview",
            "screenshot": "screenshot-preview",
        }
        blocked_states = (
            state_type.DEGRADED_READ_ONLY,
            state_type.BOOTSTRAPPING,
            state_type.SAFETY_CHECKING,
            state_type.CLOSING,
        )
        for entry in expected:
            with self.subTest(entry=entry, state="READY"):
                controller = self._read_only_controller(state_type.READY)
                context = Mock(name=f"{entry}_context")
                self.assertEqual(
                    self._invoke_read_only_controller(
                        controller, entry, context=context
                    ),
                    expected[entry],
                )
                if entry == "result":
                    controller.results.page_snapshot.assert_called_once_with(
                        limit=25, context=context
                    )
                elif entry == "private":
                    controller.results.classify_private_reference.assert_called_once_with(
                        (), 30, context=context
                    )
                elif entry == "screenshot":
                    controller.screenshots.preview.assert_called_once_with(
                        controller.paths.screenshot_inbox,
                        controller.paths.video_root,
                        context=context,
                    )
            for blocked_state in blocked_states:
                with self.subTest(entry=entry, state=blocked_state):
                    controller = self._read_only_controller(blocked_state)
                    with self.assertRaises(ControllerError):
                        self._invoke_read_only_controller(controller, entry)
                    if entry == "result":
                        controller.results.page_snapshot.assert_not_called()
                    elif entry == "task":
                        controller.task_order.list_tasks.assert_not_called()
                    elif entry == "private":
                        controller.tasks.preview.assert_not_called()
                        controller.results.classify_private_reference.assert_not_called()
                    else:
                        controller.screenshots.preview.assert_not_called()

    def test_task_scan_controller_accepts_and_forwards_operation_context(self) -> None:
        controller = self._read_only_controller(self._startup_state().READY)
        context = Mock(name="task_scan_context")

        result = controller.list_tasks(context=context)

        self.assertEqual(result, (Path("Task_A1.json"),))
        controller.task_order.list_tasks.assert_called_once_with(context=context)

    def test_runtime_diagnostic_controller_allows_only_ready_and_degraded(self) -> None:
        state_type = self._startup_state()
        for allowed_state in (state_type.READY, state_type.DEGRADED_READ_ONLY):
            with self.subTest(allowed=allowed_state):
                controller = self._runtime_diagnostic_controller(allowed_state)
                controller.health = Mock(return_value={"state": allowed_state.value})
                context = Mock(name=f"{allowed_state.value}_context")

                result = controller.runtime_status_snapshot(context=context)

                self.assertEqual(result, {"state": allowed_state.value})
                controller.health.assert_called_once_with(
                    check_processes=True, context=context
                )

        for blocked_state in (
            state_type.BOOTSTRAPPING,
            state_type.SAFETY_CHECKING,
            state_type.CLOSING,
        ):
            with self.subTest(blocked=blocked_state):
                controller = self._runtime_diagnostic_controller(blocked_state)
                controller.health = Mock()
                with self.assertRaises(ControllerError):
                    controller.runtime_status_snapshot(context=Mock())
                controller.health.assert_not_called()

    def test_runtime_diagnostic_context_cancels_between_each_access_domain(self) -> None:
        for cancel_stage in ("paths", "collector", "engine"):
            with self.subTest(cancel_stage=cancel_stage):
                controller = self._runtime_diagnostic_controller()
                context = OperationContext()
                master_is_file = controller.paths.master_settings.is_file

                if cancel_stage == "paths":
                    controller.paths.health.side_effect = lambda: (
                        context.request_cancel(),
                        {"engine_exe": True},
                    )[1]
                elif cancel_stage == "collector":
                    controller.collector.health.side_effect = lambda: (
                        context.request_cancel(),
                        False,
                    )[1]
                else:
                    controller.engine.external_running.side_effect = lambda: (
                        context.request_cancel(),
                        False,
                    )[1]

                with self.assertRaises(TaskCancelled):
                    controller.health(check_processes=True, context=context)

                if cancel_stage == "paths":
                    controller.collector.health.assert_not_called()
                    controller.engine.external_running.assert_not_called()
                elif cancel_stage == "collector":
                    controller.engine.external_running.assert_not_called()
                master_is_file.assert_not_called()

        controller = self._runtime_diagnostic_controller()
        controller.engine.external_running.return_value = True
        controller.paths.active_settings.is_file.return_value = True
        context = OperationContext()

        def cancel_during_active_settings(path: object) -> dict[str, object]:
            if path is controller.paths.active_settings:
                self.assertTrue(context.request_cancel())
                return {"run_command": "5 1 1 Q"}
            return {"accounts_urls": []}

        with patch(
            "douk_manager.controller.read_json",
            side_effect=cancel_during_active_settings,
        ):
            with self.assertRaises(TaskCancelled):
                controller.health(check_processes=True, context=context)
        controller.paths.master_settings.is_file.assert_not_called()

        controller = self._runtime_diagnostic_controller()
        controller.paths.master_settings.is_file.return_value = True
        context = OperationContext()

        def cancel_during_master(_path: Path) -> dict[str, object]:
            self.assertTrue(context.request_cancel())
            return {"accounts_urls": []}

        with patch("douk_manager.controller.read_json", side_effect=cancel_during_master):
            with self.assertRaises(TaskCancelled):
                controller.health(check_processes=True, context=context)

    def test_unknown_engine_probe_is_fail_closed_in_runtime_snapshot(self) -> None:
        engine = EngineService.__new__(EngineService)
        engine.probe_external_running = Mock(
            return_value=ProcessProbe(
                ProcessProbeState.UNKNOWN, "synthetic process probe uncertainty"
            )
        )
        self.assertTrue(engine.external_running())

        controller = self._runtime_diagnostic_controller()
        controller.engine = SimpleNamespace(
            external_running=engine.external_running,
            current=None,
        )
        snapshot = controller.health(check_processes=True)

        self.assertTrue(snapshot["engine_running"])

    def test_engine_preview_requires_ready_and_forwards_one_context(self) -> None:
        state_type = self._startup_state()
        archive = Path("synthetic-engine.zip")
        for state in (
            state_type.BOOTSTRAPPING,
            state_type.SAFETY_CHECKING,
            state_type.DEGRADED_READ_ONLY,
            state_type.CLOSING,
        ):
            with self.subTest(blocked=state), tempfile.TemporaryDirectory() as directory:
                controller = make_controller(Path(directory), engine_running=False)
                controller.startup_state = state
                with self.assertRaises(ControllerError):
                    controller.preview_engine_update(archive)

                controller.engine_updates.preview.assert_not_called()

        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=False)
            controller.startup_state = state_type.READY
            context = Mock(name="engine_preview_context")
            preview = Mock(name="engine_preview")
            controller.engine_updates.preview.return_value = preview

            result = controller.preview_engine_update(archive, context=context)

            self.assertIs(result, preview)
            controller.engine_updates.preview.assert_called_once_with(
                archive, context=context
            )

    def test_engine_apply_rejects_managed_or_healthy_collector_before_any_stage(
        self,
    ) -> None:
        for managed_running, health_running in ((True, False), (False, True)):
            with (
                self.subTest(
                    managed_running=managed_running,
                    health_running=health_running,
                ),
                tempfile.TemporaryDirectory() as directory,
            ):
                controller = make_controller(Path(directory), engine_running=False)
                controller.startup_state = self._startup_state().READY
                controller.collector.running = managed_running
                controller.collector.health.return_value = health_running
                apply_entered = threading.Event()
                backup = Mock(name="engine_update_backup")
                extract = Mock(name="engine_update_extract")
                move = Mock(name="engine_update_move")

                def unsafe_apply(*_args, **_kwargs):
                    apply_entered.set()
                    backup()
                    extract()
                    move()
                    return SimpleNamespace(
                        archive=Path(directory) / "engine.zip",
                        backup_path=Path(directory) / "backup",
                        rollback_path=Path(directory) / "rollback",
                    )

                controller.engine_updates.apply.side_effect = unsafe_apply

                error: ControllerError | None = None
                try:
                    controller.apply_engine_update(
                        Path(directory) / "engine.zip",
                        context=OperationContext(),
                    )
                except ControllerError as exc:
                    error = exc

                controller.engine_updates.apply.assert_not_called()
                self.assertFalse(apply_entered.is_set())
                backup.assert_not_called()
                extract.assert_not_called()
                move.assert_not_called()
                self.assertIsNotNone(error, "a persistent collector must reject apply")
                if not managed_running:
                    controller.collector.health.assert_called_once_with()

    def test_engine_apply_idle_collector_passes_same_context_to_service(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=False)
            controller.startup_state = self._startup_state().READY
            controller.collector.running = False
            controller.collector.health.return_value = False
            archive = Path(directory) / "engine.zip"
            context = OperationContext()
            result = SimpleNamespace(
                archive=archive,
                backup_path=Path(directory) / "backup",
                rollback_path=Path(directory) / "rollback",
            )
            controller.engine_updates.apply.return_value = result

            self.assertIs(
                controller.apply_engine_update(archive, context=context),
                result,
            )
            controller.collector.health.assert_called_once_with()
            controller.engine_updates.apply.assert_called_once_with(
                archive, context=context
            )

    def test_cancelled_screenshot_preview_preserves_files_and_never_reaches_result(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inbox, accounts, sources, destinations = self._create_screenshot_inputs(root)
            controller = self._read_only_controller(self._startup_state().READY)
            controller.paths = SimpleNamespace(
                screenshot_inbox=inbox,
                video_root=accounts,
            )
            controller.screenshots = ScreenshotService()
            context = OperationContext()
            original = sources[0].read_bytes()
            real_scan = screenshot_organizer.scan_all

            def scan_then_cancel(*args, **kwargs):
                result = real_scan(*args, **kwargs)
                self.assertTrue(context.request_cancel())
                return result

            settlements: list[tuple[object, ...]] = []
            worker = TaskWorker(
                "screenshot-preview",
                1,
                lambda worker_context: controller.screenshot_preview(
                    context=worker_context
                ),
                CancellationToken(),
                operation_context=context,
            )
            worker.settled.connect(lambda *args: settlements.append(args))
            with patch.object(
                screenshot_organizer, "scan_all", side_effect=scan_then_cancel
            ):
                worker.run()

            self.assertEqual(len(settlements), 1)
            self.assertEqual(settlements[0][:3], ("screenshot-preview", 1, TaskState.CANCELLED))
            self.assertEqual(sources[0].read_bytes(), original)
            self.assertFalse(destinations[0].exists())

    def test_degraded_read_only_rejects_every_dangerous_entry_point(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            controller = self._degraded_controller(root)
            controller.config.screenshot_post_mode = "batch"
            blocked_operations = (
                lambda: controller.update_post_options({"cleanup_after_index": True}),
                lambda: controller.create_task(
                    "A1", EarliestRule.keep(), False, "blocked", True
                ),
                lambda: controller.delete_tasks((controller.paths.tasks / "task.json",)),
                lambda: controller.generate_batches(1, 3, 1, EarliestRule.keep()),
                lambda: controller.save_task_order([]),
                lambda: controller.restore_task_order(),
                lambda: controller.activate_task(controller.paths.tasks / "task.json"),
                lambda: controller.start_current_download(),
                lambda: controller.start_monitor(),
                lambda: controller.activate_and_start(controller.paths.tasks / "task.json"),
                lambda: controller.backup_now(),
                lambda: controller.apply_engine_update(root / "engine.zip"),
                lambda: controller.start_collector(),
                lambda: controller.migrate_collector(),
                lambda: controller.organize_screenshots(),
                lambda: controller.refresh_index(),
                lambda: controller.cleanup_index(),
                lambda: controller.cleanup_index_self_test(),
                lambda: controller.run_post_actions("batch"),
            )

            for operation in blocked_operations:
                with self.subTest(operation=operation):
                    with self.assertRaises(ControllerError):
                        operation()

    def test_degraded_read_only_accepts_path_reconfiguration_without_ready_state(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            controller = self._degraded_controller(root)
            controller._build_services = Mock()
            controller.try_startup_backup = Mock(
                side_effect=AssertionError("reconfiguration must wait for a new startup result")
            )

            result = controller.reconfigure(
                {"engine_exe": str(root / "replacement" / "main.exe")}
            )

            self.assertEqual(result, "")
            self.assertNotEqual(controller.startup_state, self._startup_state().READY)
            self.assertIsNone(controller.startup_backup)
            controller._build_services.assert_called_once_with()
            controller.try_startup_backup.assert_not_called()

    def test_degraded_reconfiguration_rejects_non_path_settings(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = self._degraded_controller(Path(directory))
            controller._build_services = Mock()
            original_port = controller.config.collector_port

            with self.assertRaises(ControllerError):
                controller.reconfigure({"collector_port": 19999})

            self.assertEqual(controller.config.collector_port, original_port)
            controller._build_services.assert_not_called()
            controller.collector.stop.assert_not_called()

    def test_degraded_path_reconfiguration_does_not_stop_managed_collector(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            controller = self._degraded_controller(root)
            controller._build_services = Mock()
            controller.collector.running = True
            controller.collector.process = object()

            with self.assertRaises(ControllerError):
                controller.reconfigure({"engine_exe": str(root / "replacement" / "main.exe")})

            controller.collector.stop.assert_not_called()
            controller._build_services.assert_not_called()

    def test_degraded_path_reconfiguration_rejects_managed_engine_work(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            controller = self._degraded_controller(root)
            controller._build_services = Mock()
            managed_runtime = SimpleNamespace(running=True, mode="batch")
            controller.engine.current = managed_runtime
            original_engine_exe = controller.config.engine_exe

            with self.assertRaises(ControllerError):
                controller.reconfigure({"engine_exe": str(root / "replacement" / "main.exe")})

            self.assertIs(controller.engine.current, managed_runtime)
            self.assertEqual(controller.config.engine_exe, original_engine_exe)
            controller.engine.external_running.assert_not_called()
            controller._build_services.assert_not_called()

    def test_degraded_path_reconfiguration_skips_uncertain_process_probe(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            controller = self._degraded_controller(root)
            controller._build_services = Mock()
            controller.engine.external_running.return_value = True
            replacement = root / "replacement" / "main.exe"

            try:
                result = controller.reconfigure({"engine_exe": str(replacement)})
            except ControllerError as exc:
                self.fail(f"degraded path repair must not depend on process certainty: {exc}")

            self.assertEqual(result, "")
            self.assertEqual(controller.config.engine_exe, str(replacement))
            controller.engine.external_running.assert_not_called()
            controller._build_services.assert_called_once_with()

    def test_startup_check_is_single_flight_while_safety_checking(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=False)
            controller.startup_state = self._startup_state().BOOTSTRAPPING
            controller.startup_generation = 0

            self.assertTrue(controller.begin_startup_check(1))
            self.assertFalse(controller.begin_startup_check(2))
            self.assertEqual(controller.startup_generation, 1)
            self.assertEqual(controller.startup_state, self._startup_state().SAFETY_CHECKING)

    def test_degraded_path_repair_requires_new_generation_before_ready(self) -> None:
        assert startup_module is not None
        state = self._startup_state()
        with tempfile.TemporaryDirectory() as directory:
            controller = self._degraded_controller(Path(directory))
            controller.startup_state = state.BOOTSTRAPPING
            controller.startup_generation = 0
            controller._build_services = Mock()
            failed = startup_module.StartupSafetyResult(
                generation=4,
                success=False,
                state=state.DEGRADED_READ_ONLY,
                stage=startup_module.StartupStage.PATHS,
                summary="synthetic path failure",
                details="repair is required",
                health={},
            )
            ready = startup_module.StartupSafetyResult(
                generation=5,
                success=True,
                state=state.READY,
                stage=startup_module.StartupStage.SNAPSHOT,
                summary="ready",
                details="",
                health={},
            )

            self.assertTrue(controller.begin_startup_check(4))
            self.assertTrue(controller.apply_startup_result(failed))
            self.assertIs(controller.startup_state, state.DEGRADED_READ_ONLY)
            controller.reconfigure(
                {
                    "engine_exe": str(controller.paths.engine_exe),
                    "video_root": str(controller.paths.video_root),
                    "index_root": str(controller.paths.index_root),
                    "old_screenshot_dir": str(controller.config.old_screenshot_dir),
                }
            )

            self.assertTrue(controller.begin_startup_check(5))
            self.assertFalse(controller.apply_startup_result(failed))
            self.assertIs(controller.startup_state, state.SAFETY_CHECKING)
            self.assertTrue(controller.apply_startup_result(ready))
            self.assertIs(controller.startup_state, state.READY)

    def test_runtime_stop_and_cancel_reject_non_operational_states(self) -> None:
        state = self._startup_state()
        for blocked_state in (
            state.BOOTSTRAPPING,
            state.SAFETY_CHECKING,
            state.DEGRADED_READ_ONLY,
        ):
            with self.subTest(state=blocked_state), tempfile.TemporaryDirectory() as directory:
                controller = make_controller(Path(directory), engine_running=False)
                controller.startup_state = blocked_state
                controller.engine.current = SimpleNamespace(mode="monitor")
                controller.collector.process = object()

                for operation in (
                    controller.stop_monitor,
                    controller.cancel_current_download,
                    controller.stop_collector,
                ):
                    with self.subTest(operation=operation.__name__):
                        with self.assertRaises(ControllerError):
                            operation()

    def test_runtime_stop_and_cancel_require_manager_owned_work(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=False)
            controller.startup_state = self._startup_state().READY
            controller.engine.current = None
            controller.collector.process = None

            for operation in (
                controller.stop_monitor,
                controller.cancel_current_download,
                controller.stop_collector,
            ):
                with self.subTest(operation=operation.__name__):
                    with self.assertRaises(ControllerError):
                        operation()

    def test_runtime_stop_and_cancel_allow_managed_work_when_ready_or_closing(self) -> None:
        state = self._startup_state()
        for allowed_state in (state.READY, state.CLOSING):
            with self.subTest(state=allowed_state), tempfile.TemporaryDirectory() as directory:
                controller = make_controller(Path(directory), engine_running=False)
                controller.startup_state = allowed_state
                controller.engine.current = SimpleNamespace(mode="monitor")
                controller.engine.stop_monitor.return_value = controller.paths.active_settings
                controller.engine.cancel_batch.return_value = controller.paths.logs / "task.log"
                controller.collector.process = object()

                controller.stop_monitor()
                controller.cancel_current_download()
                controller.stop_collector()

                controller.engine.stop_monitor.assert_called_once_with()
                controller.engine.cancel_batch.assert_called_once_with(None)
                controller.collector.stop.assert_called_once_with()

    def test_monitor_start_requires_successful_startup_backup_before_engine_launch(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=False)
            controller.startup_backup = None
            controller.read_only_reason = "启动前备份失败：test failure"
            controller.try_startup_backup = Mock(return_value=controller.read_only_reason)
            controller.collector.health.return_value = False
            controller.collector.running = False
            active_before = controller.paths.active_settings.read_bytes()

            with self.assertRaisesRegex(ControllerError, "启动前备份失败"):
                controller.start_monitor()

            controller.engine.start_monitor.assert_not_called()
            self.assertEqual(controller.paths.active_settings.read_bytes(), active_before)

    def test_monitor_start_honors_existing_read_only_reason(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=False)
            controller.read_only_reason = "只读保护仍然有效"
            controller.collector.health.return_value = False
            controller.collector.running = False
            active_before = controller.paths.active_settings.read_bytes()

            with self.assertRaisesRegex(ControllerError, "只读保护"):
                controller.start_monitor()

            controller.engine.start_monitor.assert_not_called()
            self.assertEqual(controller.paths.active_settings.read_bytes(), active_before)

    def test_downloader_running_allows_collector_start_without_touching_active_or_db(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=True)
            controller.startup_backup = None
            controller.read_only_reason = "检测到下载引擎正在运行，未执行启动前备份。"
            active_before = file_sha256(controller.paths.active_settings)
            database_before = file_sha256(controller.paths.database)

            result = controller.start_collector()

            controller.collector.start.assert_called_once_with()
            controller.engine.external_running.assert_called_once_with()
            self.assertEqual(result, controller.collector.last_log_path)
            self.assertTrue(controller._last_collector_running)
            self.assertEqual(file_sha256(controller.paths.active_settings), active_before)
            self.assertEqual(file_sha256(controller.paths.database), database_before)

    def test_downloader_running_allows_collector_stop_without_touching_active_or_db(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=True)
            controller._last_collector_running = True
            active_before = file_sha256(controller.paths.active_settings)
            database_before = file_sha256(controller.paths.database)

            controller.stop_collector()

            controller.collector.stop.assert_called_once_with()
            self.assertFalse(controller._last_collector_running)
            self.assertEqual(file_sha256(controller.paths.active_settings), active_before)
            self.assertEqual(file_sha256(controller.paths.database), database_before)

    def test_collector_start_still_requires_backup_when_downloader_is_not_running(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=False)
            controller.startup_backup = None
            controller.read_only_reason = "启动前备份失败：test failure"
            controller.try_startup_backup = Mock(return_value=controller.read_only_reason)

            with self.assertRaisesRegex(ControllerError, "启动前备份失败"):
                controller.start_collector()

            controller.collector.start.assert_not_called()

    def test_downloader_running_still_rejects_dangerous_controller_operations(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=True)
            blocked_operations = (
                lambda: controller.migrate_collector(),
                lambda: controller.create_task(
                    "A1", EarliestRule.keep(), False, "blocked", True
                ),
                lambda: controller.generate_batches(1, 3, 1, EarliestRule.keep()),
                lambda: controller.activate_task(controller.paths.tasks / "task.json"),
                lambda: controller.backup_now(),
                lambda: controller.apply_engine_update(Path(directory) / "engine.zip"),
                lambda: controller.reconfigure({"engine_exe": "other-main.exe"}),
                lambda: controller.start_current_download(),
            )

            for operation in blocked_operations:
                with self.subTest(operation=operation):
                    with self.assertRaisesRegex(ControllerError, "下载引擎正在运行"):
                        operation()

            controller.collector.migrate_old_data.assert_not_called()
            controller.tasks.create_task.assert_not_called()
            controller.tasks.generate_batches.assert_not_called()
            controller.tasks.activate_existing_task.assert_not_called()
            controller.backup.create_full_snapshot.assert_not_called()
            controller.engine_updates.apply.assert_not_called()
            controller.engine.start.assert_not_called()
            self.assertFalse(hasattr(controller, "restore_backup"))

    def test_summary_lifecycle_lease_blocks_dangerous_operations_but_not_collector(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            controller = make_controller(Path(directory), engine_running=False)
            controller._download_lifecycle_active = True

            blocked_operations = (
                lambda: controller.create_task(
                    "A1", EarliestRule.keep(), False, "blocked", True
                ),
                lambda: controller.generate_batches(1, 3, 1, EarliestRule.keep()),
                lambda: controller.activate_task(controller.paths.tasks / "task.json"),
                lambda: controller.backup_now(),
                lambda: controller.apply_engine_update(Path(directory) / "engine.zip"),
                lambda: controller.reconfigure({"engine_exe": "other-main.exe"}),
                lambda: controller.update_post_options({"cleanup_after_index": True}),
                lambda: controller.start_current_download(),
            )

            for operation in blocked_operations:
                with self.subTest(operation=operation):
                    with self.assertRaisesRegex(ControllerError, "账号结果汇总"):
                        operation()

            controller.start_collector()
            controller.stop_collector()
            controller.collector.start.assert_called_once_with()
            controller.collector.stop.assert_called_once_with()

    def test_post_actions_reject_invalid_timing_before_any_sub_operation(self) -> None:
        for timing in ("none", "manual", "", "BATCH"):
            with self.subTest(timing=timing):
                controller = self._post_action_controller(
                    screenshot_mode=timing,
                    index_mode=timing,
                )

                with self.assertRaises(ControllerError):
                    controller.run_post_actions(timing)

                controller.screenshots.execute.assert_not_called()
                controller.indexer.refresh.assert_not_called()
                controller.indexer.cleanup.assert_not_called()

    def test_post_actions_reject_inactive_download_lifecycle_before_sub_operations(
        self,
    ) -> None:
        controller = self._post_action_controller(lifecycle_active=False)

        with self.assertRaises(ControllerError):
            controller.run_post_actions("batch")

        controller.screenshots.execute.assert_not_called()
        controller.indexer.refresh.assert_not_called()
        controller.indexer.cleanup.assert_not_called()

    def test_post_actions_recheck_ready_state_and_enabled_config(self) -> None:
        state = self._startup_state()
        for unavailable in (state.DEGRADED_READ_ONLY, state.CLOSING):
            with self.subTest(state=unavailable):
                controller = self._post_action_controller()
                controller.startup_state = unavailable

                with self.assertRaises(ControllerError):
                    controller.run_post_actions("batch")

                controller.screenshots.execute.assert_not_called()
                controller.indexer.refresh.assert_not_called()
                controller.indexer.cleanup.assert_not_called()

        controller = self._post_action_controller(
            screenshot_mode="none",
            index_mode="queue",
            cleanup=True,
        )

        self.assertEqual(controller.run_post_actions("batch"), [])
        controller.screenshots.execute.assert_not_called()
        controller.indexer.refresh.assert_not_called()
        controller.indexer.cleanup.assert_not_called()

    def test_post_actions_pass_one_context_through_screenshot_index_and_cleanup(
        self,
    ) -> None:
        controller = self._post_action_controller()
        context = OperationContext()
        calls: list[tuple[str, OperationContext | None]] = []

        def screenshot(*_args, context=None):
            calls.append(("screenshot", context))
            return SimpleNamespace(moved=2)

        def refresh(*_args, context=None):
            calls.append(("refresh", context))
            return controller.indexer.refresh.return_value

        def cleanup(*_args, context=None):
            calls.append(("cleanup", context))
            return controller.indexer.cleanup.return_value

        controller.screenshots.execute.side_effect = screenshot
        controller.indexer.refresh.side_effect = refresh
        controller.indexer.cleanup.side_effect = cleanup

        messages = controller.run_post_actions("batch", context=context)

        self.assertEqual(
            calls,
            [("screenshot", context), ("refresh", context), ("cleanup", context)],
        )
        self.assertEqual(
            messages,
            [
                "截图归档：2张",
                "refresh numeric summary",
                "cleanup numeric summary",
            ],
        )

    def test_post_actions_cancel_between_sub_operations_prevents_later_work(
        self,
    ) -> None:
        controller = self._post_action_controller()
        context = OperationContext()

        def screenshot(*_args, context=None):
            self.assertIs(context, context_token)
            self.assertTrue(context.request_cancel())
            return SimpleNamespace(moved=2)

        context_token = context
        controller.screenshots.execute.side_effect = screenshot

        with self.assertRaises(TaskCancelled):
            controller.run_post_actions("batch", context=context)

        controller.screenshots.execute.assert_called_once()
        controller.indexer.refresh.assert_not_called()
        controller.indexer.cleanup.assert_not_called()

    def test_cancel_winning_during_ready_gate_is_reported_as_cancelled(self) -> None:
        signature = inspect.signature(ManagerController.run_post_actions)
        self.assertIn("context", signature.parameters)

        controller = self._post_action_controller()
        context = OperationContext()
        gate_entered = threading.Event()
        release_gate = threading.Event()
        original_state = controller._current_startup_state
        errors: list[BaseException] = []

        def gated_state():
            gate_entered.set()
            if not release_gate.wait(timeout=3):
                raise AssertionError("test did not release READY gate")
            return original_state()

        controller._current_startup_state = gated_state

        def run() -> None:
            try:
                controller.run_post_actions("batch", context=context)
            except BaseException as exc:
                errors.append(exc)

        thread = threading.Thread(target=run)
        thread.start()
        self.assertTrue(gate_entered.wait(timeout=3), "READY gate was not reached")
        controller.startup_state = self._startup_state().CLOSING
        self.assertTrue(context.request_cancel())
        release_gate.set()
        thread.join(timeout=3)

        self.assertFalse(thread.is_alive())
        self.assertEqual(len(errors), 1)
        self.assertIsInstance(errors[0], TaskCancelled)
        controller.screenshots.execute.assert_not_called()
        controller.indexer.refresh.assert_not_called()
        controller.indexer.cleanup.assert_not_called()

    def test_critical_post_actions_continue_after_state_enters_closing(self) -> None:
        controller = self._post_action_controller()
        context = OperationContext()
        calls: list[str] = []

        def screenshot(*_args, context=None):
            self.assertIs(context, context_token)
            context.enter_critical_phase()
            controller.startup_state = self._startup_state().CLOSING
            self.assertFalse(context.request_cancel())
            calls.append("screenshot")
            return SimpleNamespace(moved=2)

        def refresh(*_args, context=None):
            self.assertIs(context, context_token)
            calls.append("refresh")
            return controller.indexer.refresh.return_value

        def cleanup(*_args, context=None):
            self.assertIs(context, context_token)
            calls.append("cleanup")
            return controller.indexer.cleanup.return_value

        context_token = context
        controller.screenshots.execute.side_effect = screenshot
        controller.indexer.refresh.side_effect = refresh
        controller.indexer.cleanup.side_effect = cleanup

        messages = controller.run_post_actions("batch", context=context)

        self.assertEqual(calls, ["screenshot", "refresh", "cleanup"])
        self.assertEqual(len(messages), 3)
        self.assertTrue(context.critical_to_completion)

    def test_critical_failure_in_closing_remains_failed_not_cancelled(self) -> None:
        controller = self._post_action_controller()
        context = OperationContext()

        def screenshot(*_args, context=None):
            self.assertIs(context, context_token)
            context.enter_critical_phase()
            controller.startup_state = self._startup_state().CLOSING
            return SimpleNamespace(moved=2)

        context_token = context
        controller.screenshots.execute.side_effect = screenshot
        controller.indexer.refresh.side_effect = RuntimeError("synthetic index failure")
        settlements: list[tuple[object, ...]] = []
        worker = TaskWorker(
            "post-actions",
            1,
            lambda worker_context: controller.run_post_actions(
                "batch", context=worker_context
            ),
            CancellationToken(),
            operation_context=context,
        )
        worker.settled.connect(lambda *args: settlements.append(args))

        worker.run()

        self.assertEqual(len(settlements), 1)
        self.assertIs(settlements[0][2], TaskState.FAILED)
        self.assertIsInstance(settlements[0][3], TaskFailure)
        self.assertEqual(settlements[0][3].error_type, "RuntimeError")
        self.assertIn("synthetic index failure", settlements[0][3].message)
        self.assertTrue(context.critical_to_completion)
        self.assertTrue(context.terminal_sealed)
        self.assertFalse(context.request_cancel())
        controller.indexer.cleanup.assert_not_called()

    def test_screenshot_service_cancel_during_either_scan_never_calls_safe_move(
        self,
    ) -> None:
        for blocked_scan in (1, 2):
            with self.subTest(
                blocked_scan=blocked_scan
            ), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                inbox, accounts, sources, destinations = self._create_screenshot_inputs(
                    root
                )
                context = OperationContext()
                scan_reached = threading.Event()
                release_scan = threading.Event()
                errors: list[BaseException] = []
                real_scan = screenshot_organizer.scan_all
                scan_calls = 0

                def gated_scan(*args):
                    nonlocal scan_calls
                    scan_calls += 1
                    if scan_calls == blocked_scan:
                        scan_reached.set()
                        if not release_scan.wait(timeout=3):
                            raise AssertionError("test did not release screenshot scan")
                    return real_scan(*args)

                def execute() -> None:
                    try:
                        ScreenshotService().execute(
                            inbox,
                            accounts,
                            context=context,
                        )
                    except BaseException as exc:
                        errors.append(exc)

                with (
                    patch.object(
                        screenshot_organizer,
                        "scan_all",
                        side_effect=gated_scan,
                    ),
                    patch.object(
                        screenshot_organizer,
                        "safe_move",
                        wraps=screenshot_organizer.safe_move,
                    ) as safe_move,
                ):
                    thread = threading.Thread(target=execute)
                    thread.start()
                    try:
                        self.assertTrue(
                            scan_reached.wait(timeout=3),
                            f"screenshot scan {blocked_scan} did not reach barrier",
                        )
                        self.assertTrue(context.request_cancel())
                    finally:
                        release_scan.set()
                        thread.join(timeout=3)

                self.assertFalse(thread.is_alive())
                self.assertEqual(len(errors), 1)
                self.assertIsInstance(errors[0], TaskCancelled)
                safe_move.assert_not_called()
                self.assertFalse(context.critical_to_completion)
                self.assertTrue(all(path.is_file() for path in sources))
                self.assertTrue(all(not path.exists() for path in destinations))

    def test_index_service_cancel_after_command_build_never_runs_subprocess(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source"
            source.mkdir()
            script = root / "Refresh-DoukIndex.ps1"
            script.write_text("# synthetic", encoding="ascii")
            context = OperationContext()
            command_ready = threading.Event()
            release_command = threading.Event()
            errors: list[BaseException] = []
            real_raise = context.raise_if_cancelled

            def gated_raise() -> None:
                command_ready.set()
                if not release_command.wait(timeout=3):
                    raise AssertionError("test did not release index command gate")
                real_raise()

            def refresh() -> None:
                try:
                    IndexService().refresh(
                        source,
                        root / "index",
                        root / "logs",
                        context=context,
                    )
                except BaseException as exc:
                    errors.append(exc)

            with (
                patch("douk_manager.integrations.indexer.os.name", "nt"),
                patch(
                    "douk_manager.integrations.indexer.resource_path",
                    return_value=script,
                ),
                patch.object(context, "raise_if_cancelled", side_effect=gated_raise),
                patch(
                    "douk_manager.integrations.indexer.subprocess.run"
                ) as subprocess_run,
            ):
                thread = threading.Thread(target=refresh)
                thread.start()
                try:
                    self.assertTrue(
                        command_ready.wait(timeout=3),
                        "index command did not reach cancellation gate",
                    )
                    self.assertTrue(context.request_cancel())
                finally:
                    release_command.set()
                    thread.join(timeout=3)

            self.assertFalse(thread.is_alive())
            self.assertEqual(len(errors), 1)
            self.assertIsInstance(errors[0], TaskCancelled)
            subprocess_run.assert_not_called()
            self.assertFalse(context.critical_to_completion)

    def test_empty_screenshot_then_cancel_prevents_index_critical_and_subprocess(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inbox = root / "screenshots"
            accounts = root / "accounts"
            inbox.mkdir()
            accounts.mkdir()
            controller = self._post_action_controller(cleanup=False)
            controller.paths = SimpleNamespace(
                screenshot_inbox=inbox,
                video_root=accounts,
                index_root=root / "index",
                index_refresh_logs=root / "refresh-logs",
                index_cleanup_logs=root / "cleanup-logs",
            )
            controller.screenshots = ScreenshotService()
            controller.indexer = IndexService()
            context = OperationContext()
            screenshot_finished = threading.Event()
            between_operations = threading.Event()
            release_between = threading.Event()
            errors: list[BaseException] = []
            real_execute = controller.screenshots.execute
            real_raise = context.raise_if_cancelled
            boundary_consumed = False

            def execute_screenshots(*args, **kwargs):
                result = real_execute(*args, **kwargs)
                screenshot_finished.set()
                return result

            def gated_raise() -> None:
                nonlocal boundary_consumed
                if screenshot_finished.is_set() and not boundary_consumed:
                    boundary_consumed = True
                    between_operations.set()
                    if not release_between.wait(timeout=3):
                        raise AssertionError("test did not release post-action boundary")
                real_raise()

            def run_post_actions() -> None:
                try:
                    controller.run_post_actions("batch", context=context)
                except BaseException as exc:
                    errors.append(exc)

            with (
                patch.object(
                    controller.screenshots,
                    "execute",
                    side_effect=execute_screenshots,
                ) as execute,
                patch.object(context, "raise_if_cancelled", side_effect=gated_raise),
                patch.object(
                    screenshot_organizer,
                    "safe_move",
                    wraps=screenshot_organizer.safe_move,
                ) as safe_move,
                patch(
                    "douk_manager.integrations.indexer.subprocess.run"
                ) as subprocess_run,
            ):
                thread = threading.Thread(target=run_post_actions)
                thread.start()
                try:
                    self.assertTrue(
                        between_operations.wait(timeout=3),
                        "post action did not reach screenshot/index boundary",
                    )
                    self.assertTrue(context.request_cancel())
                finally:
                    release_between.set()
                    thread.join(timeout=3)

            self.assertFalse(thread.is_alive())
            self.assertEqual(len(errors), 1)
            self.assertIsInstance(errors[0], TaskCancelled)
            execute.assert_called_once()
            safe_move.assert_not_called()
            subprocess_run.assert_not_called()
            self.assertFalse(context.critical_to_completion)

    def test_screenshot_service_critical_wins_and_completes_all_real_moves(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inbox, accounts, sources, destinations = self._create_screenshot_inputs(
                root,
                count=2,
            )
            expected_bytes = tuple(path.read_bytes() for path in sources)
            context = OperationContext()
            first_move_reached = threading.Event()
            release_first_move = threading.Event()
            errors: list[BaseException] = []
            results: list[object] = []
            real_safe_move = screenshot_organizer.safe_move
            move_calls = 0

            def gated_safe_move(plan):
                nonlocal move_calls
                move_calls += 1
                self.assertTrue(context.critical_to_completion)
                if move_calls == 1:
                    first_move_reached.set()
                    if not release_first_move.wait(timeout=3):
                        raise AssertionError("test did not release first safe move")
                return real_safe_move(plan)

            def execute() -> None:
                try:
                    results.append(
                        ScreenshotService().execute(
                            inbox,
                            accounts,
                            context=context,
                        )
                    )
                except BaseException as exc:
                    errors.append(exc)

            with patch.object(
                screenshot_organizer,
                "safe_move",
                side_effect=gated_safe_move,
            ) as safe_move:
                thread = threading.Thread(target=execute)
                thread.start()
                try:
                    self.assertTrue(
                        first_move_reached.wait(timeout=3),
                        "first safe move did not reach critical gate",
                    )
                    self.assertFalse(context.request_cancel())
                finally:
                    release_first_move.set()
                    thread.join(timeout=3)

            self.assertFalse(thread.is_alive())
            self.assertEqual(errors, [])
            self.assertEqual(len(results), 1)
            self.assertEqual(results[0].moved, 2)
            self.assertEqual(safe_move.call_count, 2)
            self.assertTrue(context.critical_to_completion)
            self.assertTrue(all(not path.exists() for path in sources))
            self.assertEqual(
                tuple(path.read_bytes() for path in destinations),
                expected_bytes,
            )

    def test_index_service_critical_wins_runs_subprocess_and_parses_summary(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source"
            source.mkdir()
            script = root / "Refresh-DoukIndex.ps1"
            script.write_text("# synthetic", encoding="ascii")
            context = OperationContext()
            process_reached = threading.Event()
            release_process = threading.Event()
            errors: list[BaseException] = []
            results: list[object] = []

            def run_subprocess(*_args, **_kwargs):
                self.assertTrue(context.critical_to_completion)
                process_reached.set()
                if not release_process.wait(timeout=3):
                    raise AssertionError("test did not release index subprocess")
                return SimpleNamespace(
                    returncode=0,
                    stdout=index_service_output("Refresh"),
                    stderr="",
                )

            def refresh() -> None:
                try:
                    results.append(
                        IndexService().refresh(
                            source,
                            root / "index",
                            root / "logs",
                            context=context,
                        )
                    )
                except BaseException as exc:
                    errors.append(exc)

            with (
                patch("douk_manager.integrations.indexer.os.name", "nt"),
                patch(
                    "douk_manager.integrations.indexer.resource_path",
                    return_value=script,
                ),
                patch(
                    "douk_manager.integrations.indexer.subprocess.run",
                    side_effect=run_subprocess,
                ) as subprocess_run,
                patch.object(
                    indexer_module,
                    "parse_index_output",
                    wraps=indexer_module.parse_index_output,
                ) as parse_output,
            ):
                thread = threading.Thread(target=refresh)
                thread.start()
                try:
                    self.assertTrue(
                        process_reached.wait(timeout=3),
                        "index subprocess did not reach critical gate",
                    )
                    self.assertFalse(context.request_cancel())
                finally:
                    release_process.set()
                    thread.join(timeout=3)

            self.assertFalse(thread.is_alive())
            self.assertEqual(errors, [])
            self.assertEqual(len(results), 1)
            self.assertEqual(results[0].output, "synthetic index details")
            self.assertEqual(results[0].summary["Created"], 5)
            subprocess_run.assert_called_once()
            parse_output.assert_called_once()
            self.assertTrue(context.critical_to_completion)

    def test_controller_real_post_services_share_context_through_all_operations(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            inbox, accounts, sources, destinations = self._create_screenshot_inputs(
                root
            )
            script = root / "index-operation.ps1"
            script.write_text("# synthetic", encoding="ascii")
            controller = self._post_action_controller()
            controller.paths = SimpleNamespace(
                screenshot_inbox=inbox,
                video_root=accounts,
                index_root=root / "index",
                index_refresh_logs=root / "refresh-logs",
                index_cleanup_logs=root / "cleanup-logs",
            )
            controller.screenshots = ScreenshotService()
            controller.indexer = IndexService()
            context = OperationContext()
            real_execute = controller.screenshots.execute
            completed = (
                SimpleNamespace(
                    returncode=0,
                    stdout=index_service_output("Refresh"),
                    stderr="",
                ),
                SimpleNamespace(
                    returncode=0,
                    stdout=index_service_output("ManualCleanup"),
                    stderr="",
                ),
            )

            def execute_then_close(*args, **kwargs):
                result = real_execute(*args, **kwargs)
                self.assertTrue(context.critical_to_completion)
                controller.startup_state = self._startup_state().CLOSING
                self.assertFalse(context.request_cancel())
                return result

            with (
                patch("douk_manager.integrations.indexer.os.name", "nt"),
                patch(
                    "douk_manager.integrations.indexer.resource_path",
                    return_value=script,
                ),
                patch(
                    "douk_manager.integrations.indexer.subprocess.run",
                    side_effect=completed,
                ) as subprocess_run,
                patch.object(
                    indexer_module,
                    "parse_index_output",
                    wraps=indexer_module.parse_index_output,
                ) as parse_output,
                patch.object(
                    controller.screenshots,
                    "execute",
                    side_effect=execute_then_close,
                ) as execute,
                patch.object(
                    controller.indexer,
                    "refresh",
                    wraps=controller.indexer.refresh,
                ) as refresh,
                patch.object(
                    controller.indexer,
                    "cleanup",
                    wraps=controller.indexer.cleanup,
                ) as cleanup,
            ):
                messages = controller.run_post_actions("batch", context=context)

            self.assertIs(execute.call_args.kwargs["context"], context)
            self.assertIs(refresh.call_args.kwargs["context"], context)
            self.assertIs(cleanup.call_args.kwargs["context"], context)
            self.assertEqual(subprocess_run.call_count, 2)
            self.assertEqual(parse_output.call_count, 2)
            self.assertEqual(len(messages), 3)
            self.assertIn("快捷方式新建5", messages[1])
            self.assertIn("重新扫描完成", messages[2])
            self.assertTrue(context.critical_to_completion)
            self.assertFalse(sources[0].exists())
            self.assertEqual(destinations[0].read_bytes(), b"screenshot-1")

    def test_post_admission_none_retires_identity_generation_binding_and_queue(
        self,
    ) -> None:
        state = self._startup_state()
        controller = SimpleNamespace(
            startup_state=state.READY,
            _download_lifecycle_active=True,
            config=SimpleNamespace(
                screenshot_post_mode="batch",
                index_post_mode="none",
                cleanup_after_index=False,
            ),
            logger=Mock(),
            release_download_lifecycle=Mock(),
        )
        controller.release_download_lifecycle.side_effect = lambda: setattr(
            controller,
            "_download_lifecycle_active",
            False,
        )
        run = SimpleNamespace(task_log=Path("logs/DownloadTask_20260820.log"))
        lifecycle_identity = "download_lifecycle:admission-none"
        window = SimpleNamespace(
            controller=controller,
            queue_output=object(),
            queue_pause_button=object(),
            queue_cancel_button=object(),
            queue_pending=[Path("A2.json")],
            queue_current=run,
            queue_active=True,
            _download_lifecycle_identity=lifecycle_identity,
            _download_post_bindings={},
            _download_post_completed_keys={"stale-lifecycle-key"},
            _background_bindings={"unrelated-task": object()},
            _background_generations={"unrelated-generation": 7},
            _record_queue_elapsed=Mock(),
            _cancel_shutdown_for_new_work=Mock(),
            _append_info=Mock(),
            coordinator=SimpleNamespace(
                start=Mock(side_effect=TaskRejectedError("synthetic conflict"))
            ),
            statusBar=Mock(
                return_value=SimpleNamespace(showMessage=Mock())
            ),
        )
        window._submit_background = Mock(
            side_effect=lambda *args, **kwargs: MainWindow._submit_background(
                window,
                *args,
                **kwargs,
            )
        )
        window._release_download_lifecycle = lambda: (
            MainWindow._release_download_lifecycle(window)
        )

        MainWindow._run_post_actions_background(window, "batch", run)

        deduplicate_key = MainWindow._download_post_action_key(
            lifecycle_identity,
            "batch",
            run,
        )
        window._submit_background.assert_called_once()
        window.coordinator.start.assert_called_once()
        self.assertNotIn(deduplicate_key, window._download_post_bindings)
        self.assertNotIn(deduplicate_key, window._background_generations)
        self.assertEqual(
            window._background_generations,
            {"unrelated-generation": 7},
        )
        self.assertEqual(set(window._background_bindings), {"unrelated-task"})
        self.assertEqual(window._download_post_completed_keys, set())
        self.assertIsNone(window._download_lifecycle_identity)
        self.assertEqual(window.queue_pending, [])
        self.assertIsNone(window.queue_current)
        self.assertFalse(window.queue_active)
        self.assertFalse(controller._download_lifecycle_active)
        controller.release_download_lifecycle.assert_called_once_with()

    def test_reconfigure_can_repair_paths_without_an_existing_startup_backup(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            controller = make_controller(root, engine_running=False)
            controller.root = controller.paths.root
            controller.config = AppConfig(
                engine_exe=str(controller.paths.engine_exe),
                video_root=str(controller.paths.video_root),
                index_root=str(controller.paths.index_root),
            )
            controller.startup_backup = None
            controller.read_only_reason = "旧路径不完整。"
            controller.collector.running = False
            controller._build_services = Mock()
            controller.try_startup_backup = Mock(return_value="新路径仍待检查。")

            result = controller.reconfigure(
                {"engine_exe": str(root / "replacement" / "main.exe")}
            )

            self.assertEqual(result, "新路径仍待检查。")
            controller.try_startup_backup.assert_called_once_with()
            controller._build_services.assert_called_once_with()

    def test_queue_options_do_not_require_a_startup_backup_while_idle(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            controller = make_controller(root, engine_running=False)
            controller.config = AppConfig(
                engine_exe=str(controller.paths.engine_exe),
                video_root=str(controller.paths.video_root),
                index_root=str(controller.paths.index_root),
            )
            controller.startup_backup = None
            controller.try_startup_backup = Mock(
                side_effect=AssertionError("queue options must not trigger a backup")
            )

            result = controller.update_post_options(
                {"cleanup_after_index": False}
            )

            self.assertTrue(result)
            self.assertFalse(controller.config.cleanup_after_index)
            controller.try_startup_backup.assert_not_called()

    def test_two_concurrent_engine_starts_launch_only_one_process(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            paths = make_test_paths(Path(directory))
            active = read_json(paths.active_settings)
            active["run_command"] = "5 1 1 Q"
            write_json_atomic(paths.active_settings, active)
            config = SimpleNamespace(batch_accounts=50, rest_seconds=150)
            services = [
                EngineService(paths, config, Mock()) for _ in range(2)
            ]
            barrier = threading.Barrier(2)
            launched = threading.Event()
            checks = [0, 0]

            def checker(index: int):
                def external_running() -> bool:
                    checks[index] += 1
                    if checks[index] == 1:
                        barrier.wait(timeout=3)
                        return False
                    return launched.is_set()

                return external_running

            for index, service in enumerate(services):
                service.external_running = Mock(side_effect=checker(index))
                service.backup.create_critical_snapshot.return_value = (
                    paths.backups / "concurrent-start"
                )

            process = Mock()
            process.poll.return_value = None
            results: list[object] = []
            errors: list[BaseException] = []

            def run(service: EngineService) -> None:
                try:
                    results.append(service.start())
                except BaseException as exc:
                    errors.append(exc)

            def launch(*args, **kwargs):
                launched.set()
                return process

            threads = [threading.Thread(target=run, args=(service,)) for service in services]
            with patch("douk_manager.core.engine.subprocess.Popen", side_effect=launch) as popen:
                for thread in threads:
                    thread.start()
                for thread in threads:
                    thread.join(timeout=5)

            self.assertTrue(all(not thread.is_alive() for thread in threads))
            self.assertEqual(len(results), 1)
            self.assertEqual(len(errors), 1)
            self.assertIsInstance(errors[0], EngineError)
            self.assertIn("已经在运行", str(errors[0]))
            self.assertEqual(popen.call_count, 1)

    @unittest.skipUnless(os.name == "nt", "Windows named mutex is required")
    def test_inherited_engine_mutex_blocks_until_child_process_exits(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            paths = make_test_paths(Path(directory))
            service = EngineService(paths, SimpleNamespace(), Mock())
            mutex = _WindowsEngineMutex.acquire(paths.engine_exe)
            child = subprocess.Popen(
                [sys.executable, "-c", "import time; time.sleep(0.4)"],
                creationflags=subprocess.CREATE_NO_WINDOW,
                **mutex.popen_options(),
            )
            mutex.close()
            try:
                with patch("douk_manager.core.engine.subprocess.run") as wmi:
                    self.assertTrue(service.external_running())
                wmi.assert_not_called()
                with self.assertRaisesRegex(EngineError, "已经在运行"):
                    _WindowsEngineMutex.acquire(paths.engine_exe)
            finally:
                child.wait(timeout=3)

            released = _WindowsEngineMutex.acquire(paths.engine_exe)
            released.close()

    @unittest.skipUnless(os.name == "nt", "Windows named mutex is required")
    def test_failed_process_launch_releases_engine_mutex(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            paths = make_test_paths(Path(directory))
            active = read_json(paths.active_settings)
            active["run_command"] = "5 1 1 Q"
            write_json_atomic(paths.active_settings, active)
            config = SimpleNamespace(batch_accounts=50, rest_seconds=150)
            service = EngineService(paths, config, Mock())
            service.external_running = Mock(return_value=False)
            service.backup.create_critical_snapshot.return_value = (
                paths.backups / "failed-start"
            )

            with patch(
                "douk_manager.core.engine.subprocess.Popen",
                side_effect=OSError("test launch failure"),
            ):
                with self.assertRaisesRegex(EngineError, "无法启动下载引擎"):
                    service.start()

            released = _WindowsEngineMutex.acquire(paths.engine_exe)
            released.close()

    def test_new_master_account_cannot_join_already_active_task_or_change_database(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            paths = make_test_paths(Path(directory), account_count=3)
            tasks = SettingsTaskService(paths, BackupService(paths))
            reusable = tasks.create_task("A1,A3", task_name="running-task")
            tasks.activate_existing_task(reusable.task_path)
            active_before = paths.active_settings.read_bytes()
            database_before = file_sha256(paths.database)

            master = read_json(paths.master_settings)
            master["accounts_urls"].append(
                {
                    "mark": "A4 newly collected",
                    "url": "https://www.douyin.com/user/new-account",
                    "tab": "post",
                    "earliest": "",
                    "latest": "",
                    "enable": True,
                }
            )
            write_json_atomic(paths.master_settings, master)

            self.assertEqual(paths.active_settings.read_bytes(), active_before)
            self.assertEqual(file_sha256(paths.database), database_before)
            with closing(
                sqlite3.connect(f"file:{paths.database.as_posix()}?mode=ro", uri=True)
            ) as connection:
                self.assertEqual(
                    connection.execute("SELECT value FROM records").fetchone()[0],
                    "safe",
                )

            tasks.activate_existing_task(reusable.task_path)
            active_after = read_json(paths.active_settings)
            self.assertEqual(
                [account["enable"] for account in active_after["accounts_urls"]],
                [True, False, True, False],
            )

    def test_real_collector_add_changes_master_and_excel_but_not_active_or_database(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            paths = make_test_paths(root, account_count=2)
            master = read_json(paths.master_settings)
            master["accounts_urls"] = [
                {
                    "mark": "A1existing111",
                    "url": "https://www.douyin.com/user/existing",
                    "tab": "post",
                    "earliest": "",
                    "latest": "",
                    "enable": True,
                },
                {
                    "mark": "",
                    "url": "",
                    "tab": "post",
                    "earliest": "",
                    "latest": "",
                    "enable": True,
                },
            ]
            write_json_atomic(paths.master_settings, master)
            create_collector_workbook(paths.collector_excel)
            active_before = file_sha256(paths.active_settings)
            database_before = file_sha256(paths.database)

            module_path = Path(collector_server.__file__).resolve()
            module_name = "douk_manager.vendor.collector_server_isolated_test"
            spec = importlib.util.spec_from_file_location(module_name, module_path)
            self.assertIsNotNone(spec)
            self.assertIsNotNone(spec.loader)
            isolated = importlib.util.module_from_spec(spec)
            environment = {
                "DOUK_COLLECTOR_DATA_DIR": str(paths.collector_data),
                "DOUK_MASTER_PATH": str(paths.master_settings),
                "DOUK_COLLECTOR_EXCEL": str(paths.collector_excel),
                "DOUK_SCREENSHOT_DIR": str(paths.screenshot_inbox),
                "DOUK_GLOBAL_LOCK_PATH": str(paths.lock_file),
            }
            with patch.dict("os.environ", environment, clear=False):
                sys.modules[module_name] = isolated
                try:
                    spec.loader.exec_module(isolated)
                finally:
                    sys.modules.pop(module_name, None)

                result = isolated.add_record(
                    {
                        "nickname": "new-account",
                        "douyin_id": "222",
                        "url": "https://www.douyin.com/user/new-user?from=search",
                    }
                )

            self.assertTrue(result["ok"])
            saved_master = read_json(paths.master_settings)
            self.assertEqual(saved_master["accounts_urls"][1]["mark"], "A2new-account222")
            self.assertEqual(
                saved_master["accounts_urls"][1]["url"],
                "https://www.douyin.com/user/new-user",
            )
            workbook = load_workbook(paths.collector_excel, keep_links=True)
            try:
                sheet = workbook[collector_server.EXCEL_SHEET_NAME]
                self.assertEqual(sheet["B3"].value, "new-account222")
                self.assertEqual(
                    sheet["D3"].value, "https://www.douyin.com/user/new-user"
                )
                self.assertIsNone(sheet["D3"].hyperlink)
            finally:
                workbook.close()
            self.assertEqual(file_sha256(paths.active_settings), active_before)
            self.assertEqual(file_sha256(paths.database), database_before)
            with closing(
                sqlite3.connect(f"file:{paths.database.as_posix()}?mode=ro", uri=True)
            ) as connection:
                self.assertEqual(
                    connection.execute("SELECT value FROM records").fetchone()[0],
                    "safe",
                )

    def test_collector_post_routes_use_shared_short_lived_file_lock(self) -> None:
        source_path = (
            Path(__file__).resolve().parents[1]
            / "src"
            / "douk_manager"
            / "vendor"
            / "collector_server.py"
        )
        tree = ast.parse(source_path.read_text(encoding="utf-8"))
        do_post = next(
            node
            for node in ast.walk(tree)
            if isinstance(node, ast.FunctionDef) and node.name == "do_POST"
        )
        lock_calls = [
            item.context_expr
            for item in ast.walk(do_post)
            if isinstance(item, ast.withitem)
            and isinstance(item.context_expr, ast.Call)
            and isinstance(item.context_expr.func, ast.Name)
            and item.context_expr.func.id == "critical_section"
        ]

        self.assertEqual(len(lock_calls), 1)
        lock_call = lock_calls[0]
        self.assertEqual(ast.unparse(lock_call.args[0]), "GLOBAL_LOCK_PATH")
        timeout = next(
            keyword.value for keyword in lock_call.keywords if keyword.arg == "timeout"
        )
        self.assertEqual(ast.literal_eval(timeout), 3.0)

    def test_collector_post_returns_busy_without_writes_when_shared_lock_is_held(
        self,
    ) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            paths = make_test_paths(root, account_count=2)
            create_collector_workbook(paths.collector_excel)
            before = {
                "master": file_sha256(paths.master_settings),
                "excel": file_sha256(paths.collector_excel),
                "active": file_sha256(paths.active_settings),
                "database": file_sha256(paths.database),
            }
            ready_path = root / "lock-ready.txt"
            lock_script = (
                "import sys,time\n"
                "from pathlib import Path\n"
                "from douk_manager.core.locks import ProcessFileLock\n"
                "with ProcessFileLock(Path(sys.argv[1])):\n"
                "    Path(sys.argv[2]).write_text('ready', encoding='ascii')\n"
                "    time.sleep(6)\n"
            )
            environment = dict(os.environ)
            environment["PYTHONDONTWRITEBYTECODE"] = "1"
            environment["PYTHONPATH"] = str(Path(__file__).resolve().parents[1] / "src")
            holder = subprocess.Popen(
                [sys.executable, "-c", lock_script, str(paths.lock_file), str(ready_path)],
                env=environment,
            )
            try:
                deadline = time.monotonic() + 3
                while not ready_path.is_file() and time.monotonic() < deadline:
                    time.sleep(0.05)
                self.assertTrue(ready_path.is_file(), "lock holder did not become ready")

                module_path = Path(collector_server.__file__).resolve()
                module_name = "douk_manager.vendor.collector_server_lock_test"
                spec = importlib.util.spec_from_file_location(module_name, module_path)
                self.assertIsNotNone(spec)
                self.assertIsNotNone(spec.loader)
                isolated = importlib.util.module_from_spec(spec)
                module_environment = {
                    "DOUK_COLLECTOR_DATA_DIR": str(paths.collector_data),
                    "DOUK_MASTER_PATH": str(paths.master_settings),
                    "DOUK_COLLECTOR_EXCEL": str(paths.collector_excel),
                    "DOUK_SCREENSHOT_DIR": str(paths.screenshot_inbox),
                    "DOUK_GLOBAL_LOCK_PATH": str(paths.lock_file),
                }
                with patch.dict("os.environ", module_environment, clear=False):
                    sys.modules[module_name] = isolated
                    try:
                        spec.loader.exec_module(isolated)
                    finally:
                        sys.modules.pop(module_name, None)

                payload = json.dumps(
                    {
                        "client_version": isolated.VERSION,
                        "page_visible": True,
                        "nickname": "blocked",
                        "douyin_id": "333",
                        "url": "https://www.douyin.com/user/blocked",
                    }
                ).encode("utf-8")
                handler = object.__new__(isolated.Handler)
                handler.path = "/add"
                handler.headers = {
                    "X-DouK-Token": isolated.ACCESS_TOKEN,
                    "Content-Length": str(len(payload)),
                }
                handler.rfile = io.BytesIO(payload)
                responses: list[tuple[dict, int]] = []
                handler._send_json = lambda body, status=200: responses.append(
                    (body, status)
                )

                handler.do_POST()

                self.assertEqual(len(responses), 1)
                self.assertEqual(responses[0][1], 409)
                self.assertEqual(responses[0][0]["code"], "MANAGER_BUSY")
                self.assertEqual(
                    {
                        "master": file_sha256(paths.master_settings),
                        "excel": file_sha256(paths.collector_excel),
                        "active": file_sha256(paths.active_settings),
                        "database": file_sha256(paths.database),
                    },
                    before,
                )
            finally:
                if holder.poll() is None:
                    holder.terminate()
                holder.wait(timeout=3)


if __name__ == "__main__":
    unittest.main()
