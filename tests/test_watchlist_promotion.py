from __future__ import annotations

import tempfile
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook, load_workbook

from douk_manager.core.collector_service import (
    FormalAddResult,
    FormalCollectorError,
    FormalInspection,
    FormalMatch,
)
from douk_manager.core.json_store import read_json
from douk_manager.core.watchlist import WatchlistService
from douk_manager.core.watchlist_promotion import (
    WatchlistPromotionError,
    WatchlistPromotionService,
)
from douk_manager.vendor import collector_server
from tests.helpers import make_test_paths


class FakeFormalService:
    def __init__(self) -> None:
        self.inspection = FormalInspection("available", (), ())
        self.add_result = FormalAddResult(1, "A1合成昵称synthetic_42", "ADDED")
        self.fail_add = False

    @staticmethod
    def payload_for_record(record: dict[str, object]) -> dict[str, object]:
        return {"url": record["url"], "nickname": record["captured_nickname"]}

    def inspect(self, _payload: dict[str, object]) -> FormalInspection:
        return self.inspection

    def preview(self, _payload: dict[str, object]):
        if self.inspection.status == "consistent":
            return type("Preview", (), {
                "status": "formal_exists",
                "next_a_number": None,
                "a_number": self.inspection.existing_a_number,
                "mark": None,
                "message_code": "FORMAL_EXISTS",
            })()
        return type("Preview", (), {
            "status": "available",
            "next_a_number": 1,
            "a_number": None,
            "mark": "A1合成昵称synthetic_42",
            "message_code": "READY",
        })()

    def add(self, _payload: dict[str, object]) -> FormalAddResult:
        if self.fail_add:
            raise FormalCollectorError("WRITE_FAILED", "合成正式写入失败。")
        return self.add_result


class WatchlistPromotionTests(unittest.TestCase):
    @staticmethod
    def _payload() -> dict[str, object]:
        return {
            "request_id": str(uuid.uuid4()),
            "url": "https://www.douyin.com/user/synthetic-42",
            "captured_nickname": "合成昵称",
            "display_name": "合成观察名称",
            "douyin_id": "synthetic_42",
            "nickname_blank": False,
            "reasons": ["few_works"],
            "note": "synthetic promotion",
            "review_after_days": 7,
        }

    @staticmethod
    def _make_excel(path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = collector_server.EXCEL_SHEET_NAME
        sheet["A2"] = "1.A1"
        sheet.merge_cells("B2:C2")
        sheet.merge_cells("D2:J2")
        workbook.save(path)
        workbook.close()

    def _fixture(self, directory: str):
        paths = make_test_paths(Path(directory), account_count=0)
        self._make_excel(paths.collector_excel)
        watchlist = WatchlistService(paths)
        watchlist.initialize(initialization_evidence=True)
        watchlist.mark_ready()
        created = watchlist.observe(self._payload())
        return paths, watchlist, created

    def test_promotion_commits_formal_pair_then_marks_w_promoted(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            paths, watchlist, created = self._fixture(directory)
            service = WatchlistPromotionService(paths, watchlist=watchlist)

            result = service.promote(created["w_id"], expected_revision=1)

            self.assertEqual(result.a_number, 1)
            self.assertFalse(result.recovered_existing)
            document = read_json(paths.watchlist)
            record = document["records"][0]
            self.assertEqual(record["state"], "promoted")
            self.assertEqual(record["promoted_a_number"], 1)
            self.assertIsNone(record["promotion_recovery"])
            self.assertEqual(document["next_w_id"], 2)
            self.assertEqual(read_json(paths.watchlist_w_watermark)["next_w_id"], 2)

            formal = read_json(paths.master_settings)["accounts_urls"][0]
            self.assertEqual(formal["mark"], "A1合成昵称synthetic_42")
            self.assertEqual(formal["url"], self._payload()["url"])
            workbook = load_workbook(paths.collector_excel, data_only=False, keep_links=True)
            sheet = workbook[collector_server.EXCEL_SHEET_NAME]
            self.assertEqual(sheet["B2"].value, "合成昵称synthetic_42")
            self.assertEqual(sheet["D2"].value, self._payload()["url"])
            self.assertIsNone(sheet["D2"].hyperlink)
            workbook.close()

    def test_formal_failure_leaves_durable_intent_and_recovery_can_complete(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            _paths, watchlist, created = self._fixture(directory)
            fake = FakeFormalService()
            fake.fail_add = True
            service = WatchlistPromotionService(
                watchlist.paths,
                watchlist=watchlist,
                formal=fake,  # type: ignore[arg-type]
            )

            with self.assertRaises(WatchlistPromotionError) as caught:
                service.promote(created["w_id"], expected_revision=1)
            self.assertEqual(caught.exception.code, "WRITE_FAILED")
            pending = read_json(watchlist.paths.watchlist)["records"][0]
            self.assertEqual(pending["state"], "watching")
            self.assertEqual(pending["promotion_recovery"]["phase"], "formal_pending")
            fake.fail_add = False
            fake.inspection = FormalInspection(
                "consistent",
                (FormalMatch(1, "", "", ("url",)),),
                (FormalMatch(1, "", "", ("url",)),),
                existing_a_number=1,
            )
            current = watchlist.snapshot()
            recovered = service.promote(
                created["w_id"],
                expected_revision=current.revision,
            )
            self.assertTrue(recovered.recovered_existing)
            self.assertEqual(recovered.a_number, 1)
            final = read_json(watchlist.paths.watchlist)["records"][0]
            self.assertEqual(final["state"], "promoted")
            self.assertEqual(final["promoted_a_number"], 1)
            self.assertIsNone(final["promotion_recovery"])

    def test_retry_is_explicit_and_uses_a_new_attempt_after_empty_reconciliation(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            _paths, watchlist, created = self._fixture(directory)
            fake = FakeFormalService()
            fake.fail_add = True
            service = WatchlistPromotionService(
                watchlist.paths,
                watchlist=watchlist,
                formal=fake,  # type: ignore[arg-type]
            )
            with self.assertRaises(WatchlistPromotionError):
                service.promote(created["w_id"], expected_revision=1)
            pending = read_json(watchlist.paths.watchlist)["records"][0]
            old_attempt = pending["promotion_recovery"]["attempt_id"]
            current = watchlist.snapshot()

            fake.fail_add = False
            result = service.retry(
                created["w_id"],
                expected_revision=current.revision,
                attempt_id=old_attempt,
            )

            self.assertEqual(result.a_number, 1)
            self.assertFalse(result.recovered_existing)
            final = read_json(watchlist.paths.watchlist)["records"][0]
            self.assertEqual(final["state"], "promoted")
            self.assertIsNone(final["promotion_recovery"])
            self.assertEqual(final["review_history"][-1]["action"], "promote")


if __name__ == "__main__":
    unittest.main()
