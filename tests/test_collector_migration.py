from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from douk_manager.config import AppConfig
from douk_manager.core.json_store import read_json, write_json_atomic
from douk_manager.integrations.collector import CollectorService
from tests.helpers import make_test_paths


class CollectorMigrationTests(unittest.TestCase):
    def test_migration_never_copies_master(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            paths = make_test_paths(root)
            master = read_json(paths.master_settings)
            master["accounts_urls"][3]["url"] = "https://www.douyin.com/user/4"
            write_json_atomic(paths.master_settings, master)
            old_root = root / "old_collector"
            old_screenshots = old_root / "账号页面截图"
            old_screenshots.mkdir(parents=True)
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Sheet1"
            for number in range(1, 9):
                row = number + 1
                sheet.merge_cells(f"B{row}:C{row}")
                sheet.merge_cells(f"D{row}:J{row}")
                sheet[f"A{row}"] = f"{number}.A{number}"
                if number <= 3:
                    sheet[f"B{row}"] = f"account{number}"
                    sheet[f"D{row}"] = (
                        "https://v.douyin.com/historical-short-link/"
                        if number == 2
                        else f"https://www.douyin.com/user/{number}"
                    )
            workbook.save(old_root / "录制名单.xlsx")
            workbook.close()
            (old_root / "settings_master.json").write_text("forbidden", encoding="utf-8")
            (old_root / "顶级.txt").write_text("1", encoding="utf-8")
            (old_screenshots / "A1.jpg").write_bytes(b"image")
            config = AppConfig(
                engine_exe=str(paths.engine_exe),
                video_root=str(paths.video_root),
                index_root=str(paths.index_root),
                old_screenshot_dir=str(old_screenshots),
            )
            service = CollectorService(config, paths)
            result = service.migrate_old_data(old_screenshots)
            self.assertIn("录制名单.xlsx", result.copied_files)
            self.assertEqual(result.copied_screenshots, 1)
            self.assertEqual(result.excel_original_max, 3)
            self.assertEqual(result.excel_final_max, 8)
            self.assertEqual(result.excel_rows_added, 5)
            self.assertEqual(result.excel_existing_url_differences, 1)
            self.assertFalse((paths.collector_data / "settings_master.json").exists())
            self.assertTrue(paths.master_settings.is_file())
            from openpyxl import load_workbook

            migrated = load_workbook(paths.collector_excel, data_only=False)
            try:
                self.assertEqual(
                    migrated["Sheet1"]["D3"].value,
                    "https://v.douyin.com/historical-short-link/",
                )
            finally:
                migrated.close()


if __name__ == "__main__":
    unittest.main()
